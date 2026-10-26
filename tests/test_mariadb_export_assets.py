import sys,unittest,json
from pathlib import Path
from concurrent.futures import ThreadPoolExecutor
from io import BytesIO
from openpyxl import load_workbook
sys.path.insert(0,str(Path(__file__).resolve().parents[2]))
from iris2.hx_configuration import DB_PATH,PDF_REPORT_DIR
from iris2.excel_export import get_heat_exchanger_export,clear_heat_exchanger_export_cache
from iris2.asset_paths import local_asset_path
from iris2.pdf_ocr_review import document_record
from iris2.database import get_engine
from sqlalchemy import text

class ExportAssetTests(unittest.TestCase):
    def test_export_has_complete_data_and_independent_streams(self):
        clear_heat_exchanger_export_cache()
        with ThreadPoolExecutor(max_workers=3) as pool:
            streams=list(pool.map(lambda _:get_heat_exchanger_export(DB_PATH),range(3)))
        payload=streams[0].getvalue()
        self.assertEqual(len({id(stream) for stream in streams}),3)
        self.assertTrue(all(stream.getvalue()==payload for stream in streams))
        streams[0].read(100)
        self.assertEqual(streams[1].tell(),0)
        workbook=load_workbook(BytesIO(payload),read_only=False,data_only=False)
        # Operational records change; compare identities with the current DB
        # rather than a historical row count (which included deleted fixtures).
        sources={"\uC5F4\uAD50\uD658\uAE30":("heat_exchanger","heat_exchanger_id"),"\uD29C\uBE0C":("heat_exchanger_tube","tube_id"),"\uAC80\uC0AC\uBCF4\uACE0\uC11C":("iris_inspection","iris_inspection_id"),"\uAC80\uC0AC\uC774\uB825":("iris_tube_result","iris_tube_result_id"),"\uC0C1\uD0DC\uBCC0\uACBD\uC774\uB825":("tube_lifecycle_event","lifecycle_event_id")}
        with get_engine().connect() as connection:
            expected={sheet:list(connection.scalars(text(f"SELECT {key} FROM {table}"))) for sheet,(table,key) in sources.items()}
        for sheet,identifiers in expected.items():
            table=next(iter(workbook[sheet].tables.values()))
            from openpyxl.utils.cell import range_boundaries
            column,start,_,end=range_boundaries(table.ref)
            exported=[workbook[sheet].cell(row,column).value for row in range(start+1,end+1)]
            self.assertCountEqual(exported,identifiers,msg=sheet)
    def test_all_referenced_pdf_files_are_local(self):
        root=Path(__file__).resolve().parents[1]
        manifest=json.loads((root/'reports/mariadb_transition/assets.json').read_text(encoding='utf-8'))
        self.assertEqual(sum(x['status']=='MISSING' for x in manifest['files']),0)
        record=document_record(1)
        path=local_asset_path(record['pdf_path'])
        self.assertTrue(path.is_relative_to(root))
        self.assertTrue(path.is_file())
        self.assertEqual(len(list(PDF_REPORT_DIR.glob('*.pdf'))),11)

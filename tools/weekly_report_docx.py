"""Build a polished IRIS weekly report DOCX from the generated report text."""

from __future__ import annotations

import argparse
import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


FONT_NAME = "Malgun Gothic"
NAVY = "17365D"
PALE_BLUE = "EAF2F8"
LIGHT_GRAY = "F5F7F9"
BORDER = "D9D9D9"


def set_run_font(run, size: float | None = None, bold: bool | None = None) -> None:
    run.font.name = FONT_NAME
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), FONT_NAME)
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), FONT_NAME)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), FONT_NAME)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    run.font.color.rgb = RGBColor(0, 0, 0)


def set_cell_shading(cell, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), fill)


def set_cell_margins(cell, top: int = 90, start: int = 110, bottom: int = 90, end: int = 110) -> None:
    properties = cell._tc.get_or_add_tcPr()
    margins = properties.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        properties.append(margins)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table) -> None:
    properties = table._tbl.tblPr
    borders = properties.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        properties.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "4")
        node.set(qn("w:color"), BORDER)


def repeat_table_header(row) -> None:
    properties = row._tr.get_or_add_trPr()
    marker = OxmlElement("w:tblHeader")
    marker.set(qn("w:val"), "true")
    properties.append(marker)


def clean_markdown(text: str) -> str:
    return text.replace("`", "").replace("**", "").strip()


def add_table(document: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    column_count = max(len(row) for row in rows)
    table = document.add_table(rows=len(rows), cols=column_count)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    set_table_borders(table)
    repeat_table_header(table.rows[0])

    for row_index, values in enumerate(rows):
        row = table.rows[row_index]
        for column_index in range(column_count):
            cell = row.cells[column_index]
            value = clean_markdown(values[column_index]) if column_index < len(values) else ""
            cell.text = value
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margins(cell)
            if row_index == 0:
                set_cell_shading(cell, NAVY)
            elif row_index % 2 == 0:
                set_cell_shading(cell, PALE_BLUE)
            else:
                set_cell_shading(cell, "FFFFFF")
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.05
            paragraph.alignment = (
                WD_ALIGN_PARAGRAPH.CENTER if column_index == column_count - 1 else WD_ALIGN_PARAGRAPH.LEFT
            )
            for run in paragraph.runs:
                set_run_font(run, 8.8, row_index == 0)
                if row_index == 0:
                    run.font.color.rgb = RGBColor(255, 255, 255)
    document.add_paragraph().paragraph_format.space_after = Pt(0)


def configure_document(document: Document) -> None:
    section = document.sections[0]
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.7)
    section.left_margin = Cm(1.9)
    section.right_margin = Cm(1.9)

    normal = document.styles["Normal"]
    normal.font.name = FONT_NAME
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_NAME)
    normal.font.size = Pt(10)
    normal.paragraph_format.space_after = Pt(4)
    normal.paragraph_format.line_spacing = 1.12

    title = document.styles["Title"]
    title.font.name = FONT_NAME
    title._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_NAME)
    title.font.size = Pt(22)
    title.font.bold = True
    title.font.color.rgb = RGBColor(0, 0, 0)

    heading = document.styles["Heading 1"]
    heading.font.name = FONT_NAME
    heading._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_NAME)
    heading.font.size = Pt(14)
    heading.font.bold = True
    heading.font.color.rgb = RGBColor(0, 0, 0)
    heading.paragraph_format.space_before = Pt(18)
    heading.paragraph_format.space_after = Pt(8)
    heading.paragraph_format.keep_with_next = True


def add_screen_samples(document: Document) -> None:
    sample_dir = Path(__file__).resolve().parent.parent / "static" / "report_samples"
    samples = [
        (sample_dir / "tube_map_2d_sample.png", "그림 1  2D Tube Map 화면 샘플"),
        (sample_dir / "tube_map_3d_sample.png", "그림 2  3D Tube Map 화면 샘플"),
    ]
    if not all(path.exists() for path, _ in samples):
        return

    heading = document.add_paragraph(style="Heading 1")
    heading.paragraph_format.page_break_before = True
    heading_run = heading.add_run("6. 2D 및 3D 화면 샘플")
    set_run_font(heading_run, 14, True)

    intro = document.add_paragraph()
    intro_run = intro.add_run(
        "실제 DB의 검사 결과를 기준으로 튜브 위치와 감육률 색상, 선택 튜브의 상세 상태를 표현한 대표 화면입니다."
    )
    set_run_font(intro_run, 9.5)
    intro.paragraph_format.space_after = Pt(6)

    for image_path, caption_text in samples:
        figure = document.add_paragraph()
        figure.alignment = WD_ALIGN_PARAGRAPH.CENTER
        figure.paragraph_format.keep_with_next = True
        figure.paragraph_format.space_after = Pt(2)
        figure.add_run().add_picture(str(image_path), width=Cm(16.5))

        caption = document.add_paragraph()
        caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
        caption.paragraph_format.space_after = Pt(7)
        caption_run = caption.add_run(caption_text)
        set_run_font(caption_run, 9, True)


def build_document(source_path: Path, output_path: Path, *, include_screen_samples: bool = True) -> None:
    document = Document()
    configure_document(document)
    source_lines = source_path.read_text(encoding="utf-8-sig").splitlines()
    index = 0
    before_first_section = True

    while index < len(source_lines):
        raw = source_lines[index].strip()
        if not raw:
            index += 1
            continue
        if raw.startswith("# "):
            paragraph = document.add_paragraph(style="Title")
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
            run = paragraph.add_run(clean_markdown(raw[2:]))
            set_run_font(run, 22, True)
            paragraph.paragraph_format.space_after = Pt(12)
            index += 1
            continue
        if raw.startswith("## "):
            before_first_section = False
            paragraph = document.add_paragraph(style="Heading 1")
            run = paragraph.add_run(clean_markdown(raw[3:]))
            set_run_font(run, 14, True)
            index += 1
            continue
        if raw.startswith("|"):
            table_rows: list[list[str]] = []
            while index < len(source_lines) and source_lines[index].strip().startswith("|"):
                cells = [cell.strip() for cell in source_lines[index].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-{3,}:?", cell) for cell in cells):
                    table_rows.append(cells)
                index += 1
            add_table(document, table_rows)
            continue
        if raw.startswith("- "):
            paragraph = document.add_paragraph(style=None)
            paragraph.paragraph_format.left_indent = Cm(0.35)
            paragraph.paragraph_format.first_line_indent = Cm(-0.25)
            paragraph.paragraph_format.space_after = Pt(3)
            bullet = paragraph.add_run("• ")
            set_run_font(bullet, 10, True)
            run = paragraph.add_run(clean_markdown(raw[2:]))
            set_run_font(run, 10)
            if before_first_section:
                paragraph.paragraph_format.left_indent = Cm(0)
                paragraph.paragraph_format.first_line_indent = Cm(0)
                bullet.text = ""
                if ":" in run.text:
                    label, value = run.text.split(":", 1)
                    run.text = ""
                    label_run = paragraph.add_run(label + ":")
                    set_run_font(label_run, 9.5, True)
                    value_run = paragraph.add_run(value)
                    set_run_font(value_run, 9.5)
            index += 1
            continue
        paragraph = document.add_paragraph()
        run = paragraph.add_run(clean_markdown(raw))
        set_run_font(run, 10)
        index += 1

    if include_screen_samples:
        add_screen_samples(document)

    properties = document.core_properties
    properties.title = "지능형열교환기 시스템 주간보고서"
    properties.subject = "프로젝트 주간 진행사항과 차주 계획"
    properties.author = "IRIS Project"
    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(output_path)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--input", required=True)
    parser.add_argument("--output", required=True)
    parser.add_argument("--no-screen-samples", action="store_true")
    args = parser.parse_args()
    build_document(
        Path(args.input),
        Path(args.output),
        include_screen_samples=not args.no_screen_samples,
    )


if __name__ == "__main__":
    main()

from __future__ import annotations

import os
import configparser
import json
import shutil
import tempfile
import time
import unittest
from datetime import datetime
from io import BytesIO
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch

from flask import render_template
from openpyxl import load_workbook


TEST_ROOT = tempfile.TemporaryDirectory()
PROJECT_DIR = Path(__file__).resolve().parent.parent
APP_DIR = PROJECT_DIR / "iris"
ORIGINAL_IRIS2_CONFIG_PATH = os.environ.get("IRIS2_CONFIG_PATH")
TEST_DB = Path(TEST_ROOT.name) / "iris_test.db"
shutil.copy2(APP_DIR / "test_data" / "heat_exchanger_v2.db", TEST_DB)
TEST_CONFIG = Path(TEST_ROOT.name) / "hx_configuration.conf"
test_configuration = configparser.ConfigParser(interpolation=None)
test_configuration.read(APP_DIR / "hx_configuration.conf", encoding="utf-8")
test_configuration.set("database", "name", str(TEST_DB))
test_configuration.set("ml", "review_database", str(Path(TEST_ROOT.name) / "ml_review_feedback.db"))
test_configuration.set("ml", "artifact_root", str(APP_DIR / "outputs" / "ml_implementation"))
test_configuration.set(
    "paths", "pdf_reports", str(APP_DIR / "iris_test_report")
)
with TEST_CONFIG.open("w", encoding="utf-8") as config_file:
    test_configuration.write(config_file)
os.environ["IRIS2_CONFIG_PATH"] = str(TEST_CONFIG)

from iris2.app import create_app
from iris2 import services
from iris2.life_prediction import (
    calculate_api510_inspection_interval,
    calculate_tube_plugging_year,
    summarize_corrosion_rates,
    summarize_remaining_life,
)
from iris2.pressure_vessel_calculations import calculate_asme_viii_internal_pressure
from iris2.iris_v2_history import (
    compare_measurement_locations,
    conservative_daily_snapshot,
    inspection_tube_length_mm,
)
from iris2.hx_configuration import DB_PATH
from iris2.iris_v2_database import session_scope
from iris2.iris_v2_models import (
    AppUser,
    Api510Calculation,
    DataQualityIssue,
    Equipment,
    HeatExchanger,
    HeatExchangerManagementCriteria,
    HeatExchangerSpecRevision,
    HeatExchangerTube,
    InspectionPlan,
    InspectionPlanEvent,
    InspectionPlanTubeTarget,
    InspectionExecutionSnapshot,
    InspectionProvider,
    IrisInspection,
    IrisIndication,
    IrisTubeResult,
    MaintenanceRecord,
    Plant,
    PressureBoundaryComponent,
    ProcessUnit,
    TubePositionEvent,
    WorkOrder,
    TubeBundle,
    HeatExchangerSideSpec,
    TubeLifecycleEvent,
    TubeLayoutPoint,
    TubeLayoutRevision,
    SystemThemePreset,
    UserScreenPermission,
    WallLossColorSetting,
)
from sqlalchemy import func, select


class IrisSystemTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.app = create_app({"TESTING": True, "SECRET_KEY": "test"})
        cls.client = cls.app.test_client()

    @classmethod
    def tearDownClass(cls):
        from iris2.excel_export import clear_heat_exchanger_export_cache
        clear_heat_exchanger_export_cache()
        TEST_ROOT.cleanup()
        if ORIGINAL_IRIS2_CONFIG_PATH is None:
            os.environ.pop("IRIS2_CONFIG_PATH", None)
        else:
            os.environ["IRIS2_CONFIG_PATH"] = ORIGINAL_IRIS2_CONFIG_PATH

    @staticmethod
    def mandatory_exchanger_form(**overrides):
        data = {
            "plant_code": "P1", "unit_code": "VCM1",
            "equipment_code": "E-MANDATORY-001",
            "equipment_tag": "US-VC-E-MANDATORY-001",
            "exchanger_code": "E-MANDATORY-001",
            "equipment_name": "Mandatory exchanger",
            "service_name": "Cooling service",
            "exchanger_type_code": "ST", "status": "OPERATING",
            "install_date": "2005-12-01",
            "shell_material": "SA-516-70", "head_material": "SA-516-70",
            "tube_material": "Cu-Ni 70/30",
            "shell_gasket_material": "Spiral wound",
            "tube_outer_diameter_mm": "19.05",
            "tube_nominal_thickness_mm": "1.25",
            "tube_length_mm": "6000", "tube_pitch_mm": "25.4",
            "pass_count": "2", "tube_count": "20",
            "shell_fluid": "Cooling Water", "head_fluid": "Cooling Water",
            "tube_fluid": "Process EDC",
            "shell_concentration_wt_percent": "0",
            "shell_material_state": "LIQUID",
            "process_criticality": "NORMAL",
            "warning_wall_loss_percent": "20",
            "critical_wall_loss_percent": "40",
        }
        data.update(overrides)
        return data

    def test_dashboard_and_primary_pages(self):
        for path in ("/", "/exchangers", "/exchangers/wall-loss-ranking", "/analysis",
                     "/settings/system-theme",
                     "/quality-issues", "/imports", "/api/health"):
            response = self.client.get(path)
            self.assertEqual(response.status_code, 200, path)
        dashboard = self.client.get("/").get_data(as_text=True)
        self.assertIn("튜브 결과", dashboard)
        self.assertIn("생산라인별·연도별 검사 현황", dashboard)
        self.assertIn("VCM1·VCM2는 생산라인 코드", dashboard)
        self.assertIn("지능형열교환기시스템", dashboard)
        self.assertIn("Inteligence Heat Exchnage System", dashboard)
        self.assertNotIn('class="brand-logo"', dashboard)
        self.assertIn("기본 설정", dashboard)
        self.assertIn("공장·생산라인", dashboard)
        self.assertIn("Wall Loss 색상", dashboard)
        self.assertIn("정밀 분석", dashboard)
        self.assertIn('href="/analysis"', dashboard)
        self.assertIn('/static/responsive-scroll.css', dashboard)
        self.assertGreaterEqual(dashboard.count('class="metric-number'), 21)
        self.assertIn('id="dashboard-metric-modal"', dashboard)
        self.assertIn('data-metric-title="열교환기"', dashboard)
        self.assertIn('data-metric-title="검사이력 보유 Tube"', dashboard)
        self.assertIn('data-metric-title="Tube 미등록 열교환기"', dashboard)
        self.assertIn('/static/hx-v6-theme.css', dashboard)
        self.assertIn('id="active-system-theme"', dashboard)
        self.assertIn('--label-color:', dashboard)
        analysis = self.client.get("/analysis").get_data(as_text=True)
        self.assertIn('class="app-shell analysis-fullscreen"', analysis)
        self.assertIn('class="precision-fullscreen-page"', analysis)
        self.assertNotIn('id="app-sidebar"', analysis)
        self.assertNotIn('class="sidebar-toggle"', analysis)
        exchanger_list = self.client.get("/exchangers").get_data(as_text=True)
        self.assertIn("울산공장", exchanger_list)
        self.assertNotIn("<th>설비코드</th>", exchanger_list)
        self.assertIn("설비번호는 상세 조회에서만 표시됩니다.", exchanger_list)
        self.assertIn('id="dashboard-metric-formula"', dashboard)
        self.assertIn("계산 로직", dashboard)
        self.assertIn("카드 숫자의 집계 기준과 의미", dashboard)
        self.assertIn('/static/dashboard-metrics.js', dashboard)
        self.assertIn('/static/dashboard-metrics.css', dashboard)
        responsive_css = self.client.get("/static/responsive-scroll.css")
        self.assertEqual(responsive_css.status_code, 200)
        self.assertIn(b"-webkit-overflow-scrolling: touch", responsive_css.data)
        self.assertIn(b"overflow-x: auto", responsive_css.data)
        responsive_css.close()
        metric_script = self.client.get("/static/dashboard-metrics.js")
        self.assertEqual(metric_script.status_code, 200)
        self.assertIn(b"data-dashboard-metric-close", metric_script.data)
        metric_script.close()
        for expected in ("VCM1", "VCM2", "2020", "2021", "445", "497", "312", "688"):
            self.assertIn(expected, dashboard)
        logo = self.client.get("/assets/hanhwa_solution2.jpg")
        self.assertEqual(logo.status_code, 200)
        self.assertEqual(logo.mimetype, "image/jpeg")
        self.assertGreater(len(logo.data), 1000)
        logo.close()
        legacy_logo = self.client.get("/assets/hanwha_solution.png")
        self.assertEqual(legacy_logo.status_code, 301)
        self.assertTrue(legacy_logo.location.endswith("/assets/hanhwa_solution2.jpg"))

    def test_system_theme_presets_can_be_saved_applied_and_reset(self):
        page = self.client.get("/settings/system-theme")
        self.assertEqual(page.status_code, 200)
        html = page.get_data(as_text=True)
        for preset_key in ("DARK", "NAVY", "CHARCOAL", "ORANGE"):
            self.assertIn(preset_key, html)

        with session_scope(DB_PATH) as session:
            presets = session.scalars(
                select(SystemThemePreset).order_by(SystemThemePreset.display_order)
            ).all()
            self.assertEqual(len(presets), 4)
            self.assertEqual(sum(bool(preset.is_active) for preset in presets), 1)

        try:
            response = self.client.post(
                "/settings/system-theme",
                data={
                    "active_preset": "NAVY",
                    "background_color_DARK": "#060e15",
                    "label_color_DARK": "#6a8494",
                    "background_color_NAVY": "#102030",
                    "label_color_NAVY": "#abcdef",
                    "background_color_CHARCOAL": "#101417",
                    "label_color_CHARCOAL": "#9aa4ab",
                    "background_color_ORANGE": "#f4f4f4",
                    "label_color_ORANGE": "#595959",
                },
                follow_redirects=True,
            )
            self.assertEqual(response.status_code, 200)
            dashboard = self.client.get("/").get_data(as_text=True)
            self.assertIn("--bg:#102030", dashboard)
            self.assertIn("--label-color:#abcdef", dashboard)
            data = services.system_theme_settings_data()
            orange_form = {"active_preset": "ORANGE"}
            for preset in data["theme_presets"]:
                for field in ("background_color", "label_color"):
                    orange_form[f"{field}_{preset['preset_key']}"] = preset[field]
            response = self.client.post("/settings/system-theme", data=orange_form, follow_redirects=True)
            self.assertEqual(response.status_code, 200)
            dashboard = self.client.get("/").get_data(as_text=True)
            self.assertIn('data-theme="orange"', dashboard)
            self.assertIn("orange-theme.css", dashboard)
            self.assertIn("--bg:#f4f4f4", dashboard)
            self.assertEqual(services.active_theme_data()["preset_key"], "ORANGE")
        finally:
            self.client.post("/settings/system-theme/reset")

        reset_dashboard = self.client.get("/").get_data(as_text=True)
        self.assertNotIn("orange-theme.css", reset_dashboard)
        self.assertIn("--bg:#060e15", reset_dashboard)
        self.assertIn("--label-color:#6a8494", reset_dashboard)

    def test_exchanger_wall_loss_ranking_orders_highest_valid_tube_result(self):
        data = services.exchanger_wall_loss_ranking()
        values = [
            row["wall_loss_percent"] for row in data["rows"]
            if row["wall_loss_percent"] is not None
        ]
        self.assertTrue(values)
        self.assertEqual(values, sorted(values, reverse=True))
        self.assertEqual(
            len(data["rows"]), services.dashboard_data()["counts"]["exchangers"]
        )
        self.assertEqual(sum(data["status_counts"].values()), len(data["rows"]))
        first = data["rows"][0]
        filtered = services.exchanger_wall_loss_ranking(
            search=first["equipment_code"]
        )
        self.assertTrue(filtered["rows"])
        self.assertTrue(all(
            row["equipment_code"] == first["equipment_code"]
            for row in filtered["rows"]
        ))
        page = self.client.get("/exchangers/wall-loss-ranking")
        self.assertEqual(page.status_code, 200)
        html = page.get_data(as_text=True)
        self.assertIn("최대 감육률 순위", html)
        self.assertIn("동일값은 최근 검사 우선", html)
        self.assertIn(first["equipment_code"], html)

    def test_tube_layout_revision_management_is_read_only(self):
        response = self.client.get("/settings/tube-layout-revisions")
        self.assertEqual(response.status_code, 200)
        page = response.get_data(as_text=True)
        self.assertIn("Tube Layout Revision 관리", page)
        self.assertIn("읽기 전용", page)
        self.assertIn("Physical / Layout", page)
        filtered = self.client.get(
            "/settings/tube-layout-revisions?status=APPROVED&equipment=E-"
        )
        self.assertEqual(filtered.status_code, 200)

    def test_inspection_provider_software_management(self):
        with session_scope(DB_PATH) as session:
            inspection_id = session.scalar(
                select(IrisInspection.iris_inspection_id).order_by(
                    IrisInspection.iris_inspection_id
                )
            )
        response = self.client.post(
            f"/inspections/{inspection_id}/execution",
            data={
                "new_provider_code": "TEST_NDT",
                "new_provider_name": "Test NDT Company",
                "inspection_company_raw": "Test NDT Co.",
                "contract_no": "CT-2026-01",
                "procedure_no": "IRIS-P-01",
                "procedure_revision": "2",
                "acquisition_software_name": "AcquireX",
                "acquisition_software_version": "5.1",
                "analysis_software_name": "AnalyzeX",
                "analysis_software_version": "3.2",
                "equipment_calibration_date": "2026-01-02",
                "equipment_calibration_valid_until": "2027-01-01",
                "organization_examiner": "Test NDT Company",
            },
            follow_redirects=True,
        )
        self.assertEqual(response.status_code, 200)
        page = response.get_data(as_text=True)
        self.assertIn("Test NDT Company", page)
        self.assertIn("AcquireX", page)
        self.assertIn('class="governance-form"', page)
        self.assertIn("계약·검사 절차", page)
        self.assertIn("검사 장비·교정", page)
        self.assertIn("탐촉자·담당자", page)
        self.assertIn("inspection-detail.css", page)
        with session_scope(DB_PATH) as session:
            provider = session.scalar(select(InspectionProvider).where(
                InspectionProvider.provider_code == "TEST_NDT"
            ))
            execution = session.scalar(select(InspectionExecutionSnapshot).where(
                InspectionExecutionSnapshot.iris_inspection_id == inspection_id
            ))
            self.assertIsNotNone(provider)
            self.assertEqual(execution.inspection_provider_id, provider.inspection_provider_id)
            self.assertEqual(execution.analysis_software_version, "3.2")

    def test_normalized_location_is_always_displayed_in_millimetres(self):
        with session_scope(DB_PATH) as session:
            result = session.scalar(
                select(IrisTubeResult)
                .where(IrisTubeResult.location_mm.is_not(None))
                .order_by(IrisTubeResult.iris_tube_result_id)
            )
            self.assertIsNotNone(result)
            result_id = result.iris_tube_result_id
            inspection_id = result.iris_inspection_id
            original_location = result.location_mm
            original_unit = result.location_source_unit
            result.location_mm = 3968.0
            result.location_source_unit = "m"
        try:
            page = self.client.get(f"/inspections/{inspection_id}")
            self.assertEqual(page.status_code, 200)
            html = page.get_data(as_text=True)
            self.assertIn("3968.0 mm", html)
            self.assertNotIn("3968.0 m</td>", html)
        finally:
            with session_scope(DB_PATH) as session:
                saved = session.get(IrisTubeResult, result_id)
                saved.location_mm = original_location
                saved.location_source_unit = original_unit

    def test_quality_issues_is_last_basic_settings_item(self):
        dashboard = self.client.get("/").get_data(as_text=True)
        settings_start = dashboard.index('class="sidebar-group-menu settings-menu"')
        settings_end = dashboard.index("</details>", settings_start)
        settings_menu = dashboard[settings_start:settings_end]
        self.assertEqual(dashboard.count('href="/quality-issues"'), 1)
        self.assertGreater(
            settings_menu.index("품질 이슈"),
            settings_menu.index("Wall Loss 색상"),
        )

        quality_page = self.client.get("/quality-issues").get_data(as_text=True)
        quality_settings_start = quality_page.index(
            'class="sidebar-group-menu settings-menu"'
        )
        self.assertIn(
            "open",
            quality_page[quality_settings_start:quality_settings_start + 100],
        )

    def test_quality_issue_can_be_resolved_with_audit_note(self):
        with session_scope(DB_PATH) as session:
            issue_id = session.scalar(
                select(DataQualityIssue.data_quality_issue_id)
                .where(DataQualityIssue.status == "OPEN")
                .order_by(DataQualityIssue.data_quality_issue_id)
            )
        if issue_id is None:
            self.skipTest("품질 이슈 fixture가 없습니다.")
        response = self.client.post(
            f"/quality-issues/{issue_id}/resolve",
            data={"resolution_note": "원천 대조 후 정상값으로 확인", "return_status": "OPEN"},
            follow_redirects=True,
        )
        self.assertEqual(response.status_code, 200)
        with session_scope(DB_PATH) as session:
            issue = session.get(DataQualityIssue, issue_id)
            self.assertEqual(issue.status, "RESOLVED")
            self.assertEqual(issue.resolution_note, "원천 대조 후 정상값으로 확인")
            self.assertIsNotNone(issue.reviewed_at)

    def test_exchanger_list_supports_safe_sorting(self):
        attributes = {
            "plant": (3, "plant_code"),
            "unit": (2, "unit_code"),
            "equipment": (1, "equipment_code"),
        }
        for sort_by, (row_index, attribute) in attributes.items():
            ascending = services.list_exchangers(sort_by=sort_by, sort_dir="asc")["rows"]
            descending = services.list_exchangers(sort_by=sort_by, sort_dir="desc")["rows"]
            ascending_values = [getattr(row[row_index], attribute) for row in ascending]
            descending_values = [getattr(row[row_index], attribute) for row in descending]
            self.assertEqual(ascending_values, sorted(ascending_values))
            self.assertEqual(descending_values, sorted(descending_values, reverse=True))

        page = self.client.get("/exchangers?sort_by=equipment&sort_dir=desc")
        self.assertEqual(page.status_code, 200)
        self.assertIn(b'name="sort_by"', page.data)
        self.assertIn(b'value="equipment" selected', page.data)
        self.assertIn(b'name="sort_dir"', page.data)
        self.assertIn(b'value="desc" selected', page.data)

        fallback = services.list_exchangers(sort_by="unsafe SQL", sort_dir="sideways")
        self.assertEqual(fallback["sort_by"], "plant")
        self.assertEqual(fallback["sort_dir"], "asc")

    def test_exchanger_code_search_survives_numeric_equipment_code_change(self):
        with session_scope(DB_PATH) as session:
            equipment = session.scalar(select(Equipment).where(Equipment.equipment_id == 1))
            spec = session.scalar(select(HeatExchangerSpecRevision).where(
                HeatExchangerSpecRevision.heat_exchanger_id == 1,
                HeatExchangerSpecRevision.is_current.is_(True),
            ))
            old_equipment_code = equipment.equipment_code
            exchanger_code = spec.exchanger_code
            equipment.equipment_code = "9900001"
        try:
            data = services.list_exchangers(search=exchanger_code)
            self.assertTrue(any(row[0].heat_exchanger_id == 1 for row in data["rows"]))
            response = self.client.get(f"/api/v1/exchangers?search={exchanger_code}")
            self.assertEqual(response.status_code, 200)
            row = next(item for item in response.get_json() if item["heat_exchanger_id"] == 1)
            self.assertEqual(row["equipment_code"], "9900001")
            self.assertEqual(row["exchanger_code"], exchanger_code)
        finally:
            with session_scope(DB_PATH) as session:
                session.get(Equipment, 1).equipment_code = old_equipment_code

    def test_qr_exchanger_lookup_renders_compact_authenticated_view(self):
        response = self.client.get("/q/exchangers/1")
        self.assertEqual(response.status_code, 200)
        self.assertIn("QR 전용 간편조회 화면".encode("utf-8"), response.data)
        self.assertIn(b"/exchangers/1/tube-map", response.data)
        self.assertIn("검사 이력·PDF".encode("utf-8"), response.data)
        self.assertIn("최신 검사 PDF".encode("utf-8"), response.data)
        self.assertIn(b"/inspections/", response.data)
        self.assertIn(b"/pdf", response.data)

    def test_authentication_and_screen_permissions(self):
        with session_scope(DB_PATH) as db_session:
            db_session.query(UserScreenPermission).delete()
            db_session.query(AppUser).filter(
                AppUser.is_active.is_(True)
            ).delete(synchronize_session=False)
        auth_app = create_app({"TESTING": True, "SECRET_KEY": "auth-test",
                               "AUTH_DISABLED": False})
        client = auth_app.test_client()
        first = client.get("/")
        self.assertEqual(first.status_code, 302)
        self.assertTrue(first.location.endswith("/setup-admin"))
        invalid_email = client.post("/setup-admin", data={
            "username": "admin", "display_name": "시스템 관리자",
            "password": "abcdefghij", "password_confirm": "abcdefghij",
        })
        self.assertEqual(invalid_email.status_code, 200)
        self.assertIn("올바른 이메일 주소".encode("utf-8"), invalid_email.data)
        invalid_password = client.post("/setup-admin", data={
            "username": "admin@example.com", "display_name": "시스템 관리자",
            "password": "abcdefgh", "password_confirm": "abcdefgh",
        })
        self.assertEqual(invalid_password.status_code, 200)
        self.assertIn("9자 이상".encode("utf-8"), invalid_password.data)
        created = client.post("/setup-admin", data={
            "username": "admin@example.com", "display_name": "시스템 관리자",
            "password": "!@#$%^&*()", "password_confirm": "!@#$%^&*()",
        }, follow_redirects=True)
        self.assertEqual(created.status_code, 200)
        self.assertIn("초기 관리자 계정이 생성되었습니다.".encode("utf-8"), created.data)
        users_page = client.get("/settings/users")
        self.assertEqual(users_page.status_code, 200)
        database_monitor_page = client.get("/settings/database-monitor")
        self.assertEqual(database_monitor_page.status_code, 200)
        self.assertIn("DB 상태 점검".encode("utf-8"), database_monitor_page.data)
        self.assertIn(b"PRAGMA integrity_check", database_monitor_page.data)
        backup_manager_page = client.get("/settings/backups")
        self.assertEqual(backup_manager_page.status_code, 200)
        self.assertIn("백업 관리".encode("utf-8"), backup_manager_page.data)
        self.assertIn("backup_settings.json".encode("utf-8"), backup_manager_page.data)
        self.assertIn("화면별 권한".encode("utf-8"), users_page.data)
        self.assertIn(b'/static/auth.css', users_page.data)
        viewer = client.post("/settings/users", data={
            "username": "viewer1@example.com", "display_name": "조회 사용자",
            "password": "abcdefghij", "password_confirm": "abcdefghij",
            "is_active": "1",
            "permission_DASHBOARD_view": "1",
            "permission_EXCHANGERS_view": "1",
        }, follow_redirects=True)
        self.assertIn("사용자가 등록되었습니다.".encode("utf-8"), viewer.data)
        client.post("/logout")
        qr_scan = client.get("/q/exchangers/1")
        self.assertEqual(qr_scan.status_code, 302)
        self.assertIn("/login?next=/q/exchangers/1", qr_scan.location)
        logged_in = client.post("/login", data={
            "username": "VIEWER1@example.com", "password": "abcdefghij",
            "next": "/q/exchangers/1",
        })
        self.assertEqual(logged_in.status_code, 302)
        self.assertTrue(logged_in.location.endswith("/q/exchangers/1"))
        qr_after_login = client.get(logged_in.location)
        self.assertEqual(qr_after_login.status_code, 200)
        self.assertIn("QR 전용 간편조회 화면".encode("utf-8"), qr_after_login.data)
        logged_in = client.get("/")
        current_user_html = logged_in.get_data(as_text=True).split(
            '<div class="current-user">', 1
        )[1].split('</div>', 1)[0]
        self.assertIn("viewer1@example.com", current_user_html)
        self.assertIn(">사용자</button>", current_user_html)
        self.assertNotIn("조회 사용자", current_user_html)
        self.assertNotIn(">로그아웃</button>", current_user_html)
        exchanger_page = client.get("/exchangers")
        self.assertEqual(exchanger_page.status_code, 200)
        self.assertIn("내 기본 설정".encode("utf-8"), exchanger_page.data)
        preferences_page = client.get("/settings/preferences")
        self.assertEqual(preferences_page.status_code, 200)
        self.assertIn("검사 PDF 열기 방식".encode("utf-8"), preferences_page.data)
        saved_preferences = client.post(
            "/settings/preferences",
            data={"pdf_open_mode": "CURRENT"},
            follow_redirects=True,
        )
        self.assertIn("사용자별 기본설정이 저장되었습니다.".encode("utf-8"), saved_preferences.data)
        exchanger_page = client.get("/exchangers")
        pdf_link_tag = exchanger_page.get_data(as_text=True).split(
            'class="button report-view"', 1
        )[1].split(">", 1)[0]
        self.assertNotIn('target="_blank"', pdf_link_tag)
        with session_scope(DB_PATH) as db_session:
            viewer_user = db_session.scalar(
                select(AppUser).where(AppUser.username == "viewer1@example.com")
            )
            self.assertEqual(viewer_user.pdf_open_mode, "CURRENT")
        self.assertEqual(client.get("/api/v1/exchangers").status_code, 200)
        self.assertEqual(client.get("/inspections/1/pdf").status_code, 403)
        self.assertNotEqual(client.get("/api/exchangers/1/tube-map").status_code, 403)
        self.assertEqual(
            client.post("/api/tubes/1/position", json={"row_no": 1, "column_no": 1}).status_code,
            403,
        )
        vendor_response = client.get("/assets/three.min.js")
        try:
            self.assertNotEqual(vendor_response.status_code, 403)
        finally:
            vendor_response.close()
        nav = exchanger_page.get_data(as_text=True)
        nav = nav[nav.index("<nav>"):nav.index("</nav>")]
        self.assertIn("열교환기 목록", nav)
        self.assertNotIn("열교환기 입력", nav)
        self.assertNotIn("정비 관리", nav)
        self.assertEqual(client.get("/exchangers/new").status_code, 403)
        self.assertEqual(client.get("/maintenance").status_code, 403)
        self.assertEqual(client.get("/settings/users").status_code, 403)
        self.assertEqual(client.get("/settings/database-monitor").status_code, 403)
        self.assertEqual(client.get("/settings/backups").status_code, 403)
        with client.session_transaction() as auth_session:
            auth_session["last_activity_at"] = time.time() - (30 * 60 + 1)
        expired = client.get("/exchangers")
        self.assertEqual(expired.status_code, 302)
        self.assertIn("/login?next=/exchangers", expired.location)
        with client.session_transaction() as auth_session:
            self.assertNotIn("user_id", auth_session)
        api_expired = client.get("/api/v1/exchangers")
        self.assertEqual(api_expired.status_code, 302)
        self.assertIn("/login?next=/api/v1/exchangers", api_expired.location)

    def test_csrf_is_required_for_cookie_authenticated_forms(self):
        csrf_app = create_app({
            "TESTING": True, "AUTH_DISABLED": True, "CSRF_ENABLED": True,
            "SECRET_KEY": "csrf-test-secret-key-at-least-32-characters",
        })
        client = csrf_app.test_client()
        target = "/login" if services.user_count() else "/setup-admin"
        client.get(target)
        self.assertEqual(client.post(target, data={}).status_code, 400)
        with client.session_transaction() as csrf_session:
            token = csrf_session["csrf_token"]
        response = client.post(target, data={"csrf_token": token})
        self.assertNotEqual(response.status_code, 400)

    def test_existing_exchanger_inspection_and_tube_history(self):
        exchanger_list = self.client.get("/exchangers")
        self.assertEqual(exchanger_list.status_code, 200)
        html = exchanger_list.get_data(as_text=True)
        nav_html = html[html.index("<nav>"):html.index("</nav>")]
        self.assertIn("열교환기 관리", nav_html)
        self.assertNotIn("정밀 분석", nav_html)
        self.assertIn("열교환기 목록", nav_html)
        self.assertIn("열교환기 입력", nav_html)
        self.assertIn("JSON 로드", nav_html)
        self.assertNotIn("Tube Plugging 예측", nav_html)
        self.assertIn("Excel Export", nav_html)
        self.assertEqual(nav_html.count('class="sidebar-group-menu'), 3)
        self.assertEqual(nav_html.count('class="sidebar-submenu"'), 3)
        self.assertIn('href="/inspections/tubes"', nav_html)
        self.assertIn("튜브 검사이력", nav_html)
        self.assertIn('href="/exchangers/new"', nav_html)
        self.assertIn('href="/imports"', nav_html)
        self.assertIn('href="/exports/heat-exchangers-tubes.xlsx"', nav_html)
        topbar_html = html[html.index('<header class="topbar">'):html.index("</header>")]
        self.assertIn('class="top-analysis-tab', topbar_html)
        self.assertIn("정밀 분석", topbar_html)
        self.assertIn('href="/analysis"', topbar_html)
        self.assertIn("JSON 적재", topbar_html)

        analysis = self.client.get("/analysis")
        self.assertEqual(analysis.status_code, 200)
        analysis_html = analysis.get_data(as_text=True)
        self.assertIn("Tube Sheet Map", analysis_html)
        self.assertIn("결함 목록", analysis_html)
        self.assertIn("잔여수명 계산 방법", analysis_html)
        self.assertIn("열교환기 등록", topbar_html)
        self.assertLess(topbar_html.index("JSON 적재"), topbar_html.index("열교환기 등록"))
        detail = self.client.get("/exchangers/1")
        self.assertEqual(detail.status_code, 200)
        self.assertIn("Tube Plugging 예측".encode(), detail.data)
        self.assertIn("튜브별 Plugging 도달 연도".encode(), detail.data)
        self.assertIn("최근 검사일 +", detail.get_data(as_text=True))
        self.assertIn("연도별 Tube Plugging 누적 예측".encode(), detail.data)
        self.assertIn("결함 목록".encode(), detail.data)
        self.assertIn(b"life-bar-chart", detail.data)
        self.assertIn(b"life-prediction.js", detail.data)
        self.assertIn(b"equipment-summary-strip", detail.data)
        self.assertIn("설비코드".encode("utf-8"), detail.data)
        self.assertIn("검사수량 / 총수량".encode("utf-8"), detail.data)
        self.assertIn("부식속도 (mm/yr)".encode("utf-8"), detail.data)
        self.assertEqual(self.client.get("/inspections/1").status_code, 200)
        with session_scope(DB_PATH) as session:
            tube_id = session.scalar(select(HeatExchangerTube.tube_id).limit(1))
        self.assertEqual(self.client.get(f"/tubes/{tube_id}").status_code, 200)

    def test_dashboard_shows_exchanger_and_tube_data_entry_coverage(self):
        coverage = services.dashboard_data()["data_entry_coverage"]
        self.assertEqual(coverage["exchanger"]["field_count"], 26)
        self.assertEqual(coverage["exchanger"]["total_count"],
                         services.dashboard_data()["counts"]["exchangers"])
        self.assertEqual(coverage["tube"]["field_count"], 4)
        self.assertEqual(coverage["tube"]["total_count"],
                         services.dashboard_data()["counts"]["tubes"])
        counts = services.dashboard_data()["counts"]
        self.assertGreater(counts["tubes"], counts["registered_tubes"])
        self.assertEqual(
            counts["exchangers_without_registered_tubes"],
            counts["exchangers"] - counts["exchangers_with_registered_tubes"],
        )
        self.assertEqual(
            counts["exchangers_without_registered_tubes"],
            counts["tube_unspecified_without_registered_tubes"]
            + counts["tube_specified_without_registered_tubes"],
        )
        self.assertEqual(
            coverage["tube_population"]["total_count"], counts["tubes"]
        )
        inspection_history = next(
            group for group in coverage["tube_groups"]
            if group["label"] == "\uac80\uc0ac\uc774\ub825 \ubcf4\uc720 Tube"
        )
        self.assertEqual(
            inspection_history["note"],
            "사양상 전체 Tube 수 대비 연도와 관계없이 검사결과가 1건 이상 연결된 Tube",
        )
        thickness_history = next(
            group for group in coverage["tube_groups"]
            if group["label"] == "\ucd5c\uc18c\ub450\uaed8 \uce21\uc815\uc774\ub825 \ubcf4\uc720 Tube"
        )
        self.assertEqual(
            thickness_history["note"],
            "사양상 전체 Tube 수 대비 최소 잔존두께가 1건 이상 등록된 Tube",
        )
        for summary in (coverage["exchanger"], coverage["tube"],
                        coverage["exchanger_management"]):
            self.assertGreaterEqual(summary["rate"], 0)
            self.assertLessEqual(summary["rate"], 100)

        response = self.client.get("/")
        self.assertEqual(response.status_code, 200)
        html = response.get_data(as_text=True)
        self.assertIn("열교환기 주요 데이터 입력률", html)
        self.assertIn("Tube 주요 데이터 입력률", html)
        self.assertIn("Tube별 유체통로(Pass)", html)
        self.assertIn("개별 Tube가 속한 유체 흐름 구간", html)
        self.assertIn("최소허용두께와 검사 주기 기준", html)
        self.assertIn("QR-code 읽기", html)
        self.assertIn("열교환기 ID 숫자만 저장된 QR-code", html)
        self.assertIn('placeholder="예: 176"', html)
        self.assertIn('data-qr-scan-open', html)
        self.assertIn('/static/vendor/jsqr/jsQR.js', html)
        self.assertIn('/static/dashboard-qr-scanner.js', html)
        vendor_response = self.client.get('/static/vendor/jsqr/jsQR.js')
        try:
            self.assertEqual(vendor_response.status_code, 200)
        finally:
            vendor_response.close()
        self.assertIn("동일 위치 재검사", html)
        self.assertIn("연속 검사일 비교쌍", html)
        self.assertIn("전체 튜브", html)
        self.assertIn("개별 등록", html)
        self.assertIn('aria-label="주요 데이터 입력률"', html)

    def test_data_entry_group_counts_slots_and_complete_records(self):
        summary = services._data_entry_group(
            "test",
            [{"a": "value", "b": None}, {"a": "value", "b": 0}],
            ("a", "b"),
        )
        self.assertEqual(summary["populated_count"], 3)
        self.assertEqual(summary["possible_count"], 4)
        self.assertEqual(summary["complete_count"], 1)
        self.assertEqual(summary["rate"], 75.0)

    def test_same_position_reinspection_coverage_uses_position_matches(self):
        matched = json.dumps({
            "thickness_position_comparison": {"matched": True},
            "wall_loss_position_comparison": {"matched": False},
        })
        unmatched = json.dumps({
            "thickness_position_comparison": {"matched": False},
            "wall_loss_position_comparison": {"matched": False},
        })
        matched_by_loss = json.dumps({
            "wall_loss_position_comparison": {"matched": True},
        })
        summary = services._same_position_reinspection_coverage(
            [matched, unmatched, matched_by_loss]
        )
        self.assertEqual(summary["complete_count"], 2)
        self.assertEqual(summary["total_count"], 3)
        self.assertEqual(summary["rate"], 66.7)

    def test_tube_inspection_history_filters_and_paginates(self):
        data = services.tube_inspection_history_data()
        self.assertGreater(data["total_count"], 0)
        self.assertLessEqual(len(data["rows"]), data["per_page"])
        self.assertEqual(
            data["page_count"],
            max(1, (data["total_count"] + data["per_page"] - 1) // data["per_page"]),
        )
        sample = data["rows"][0]
        response = self.client.get("/inspections/tubes", query_string={
            "exchanger_id": sample["heat_exchanger_id"],
            "year": sample["inspection_year"],
            "tube_no": sample["tube"].tube_no,
        })
        self.assertEqual(response.status_code, 200)
        self.assertIn("열교환기·연도·튜브별 검사이력".encode("utf-8"), response.data)
        self.assertIn(b"Inspection summary", response.data)
        self.assertIn("2D Tube Map 보기".encode("utf-8"), response.data)
        self.assertIn(
            f'href="/exchangers/{sample["heat_exchanger_id"]}/tube-map"'.encode(),
            response.data,
        )
        self.assertIn(sample["exchanger_code"].encode("utf-8"), response.data)
        self.assertIn(sample["tube"].tube_no.encode("utf-8"), response.data)
        self.assertIn(str(sample["inspection_year"]).encode(), response.data)
        self.assertIn(b'href="/tubes/', response.data)
        self.assertIn(b'href="/inspections/', response.data)
        exchanger_asc = services.tube_inspection_history_data(
            sort_by="exchanger", sort_dir="asc"
        )
        exchanger_codes = [row["exchanger_code"] for row in exchanger_asc["rows"]]
        self.assertEqual(exchanger_codes, sorted(exchanger_codes))
        tube_desc = services.tube_inspection_history_data(
            sort_by="tube", sort_dir="desc"
        )
        tube_keys = [(
            row["tube"].row_no if row["tube"].row_no is not None
            else row["result"].source_row if row["result"].source_row is not None else -1,
            row["tube"].column_no if row["tube"].column_no is not None
            else row["result"].source_column if row["result"].source_column is not None else -1,
            row["tube"].tube_no,
        ) for row in tube_desc["rows"]]
        self.assertEqual(tube_keys, sorted(tube_keys, reverse=True))
        sorted_page = self.client.get(
            "/inspections/tubes?sort_by=tube&sort_dir=desc"
        )
        self.assertIn(b'<option value="tube" selected>', sorted_page.data)
        self.assertIn(b'<option value="desc" selected>', sorted_page.data)
        self.assertIn(b"sort_by=tube", sorted_page.data)
        self.assertIn(b"sort_dir=desc", sorted_page.data)
        self.assertEqual(
            self.client.get("/inspections/tubes?year=bad&page=bad").status_code,
            200,
        )
        invalid_sort = services.tube_inspection_history_data(
            sort_by="not-allowed", sort_dir="sideways"
        )
        self.assertEqual(invalid_sort["sort_by"], "exchanger")
        self.assertEqual(invalid_sort["sort_dir"], "asc")

    def test_ml_review_targets_are_read_only_and_fail_closed_for_stale_snapshot(self):
        before = DB_PATH.read_bytes()
        response = self.client.get("/inspections/ml-review")
        self.assertEqual(response.status_code, 200)
        self.assertIn("ML 이상치 검토 우선순위".encode("utf-8"), response.data)
        self.assertIn("검토지원 비활성".encode("utf-8"), response.data)
        self.assertIn(b'href="/inspections/ml-review"', response.data)
        self.assertEqual(DB_PATH.read_bytes(), before)

    def test_ml_spatial_hotspots_are_read_only_and_fail_closed_for_stale_snapshot(self):
        before = DB_PATH.read_bytes()
        response = self.client.get("/inspections/ml-hotspots")
        self.assertEqual(response.status_code, 200)
        self.assertIn("공간 부식 Hotspot 검토".encode("utf-8"), response.data)
        self.assertIn("공간분석 비활성".encode("utf-8"), response.data)
        self.assertIn(b'href="/inspections/ml-hotspots"', response.data)
        self.assertEqual(DB_PATH.read_bytes(), before)

    def test_ml_degradation_patterns_are_read_only_and_fail_closed_for_stale_snapshot(self):
        before = DB_PATH.read_bytes()
        response = self.client.get("/inspections/ml-patterns")
        self.assertEqual(response.status_code, 200)
        self.assertIn("유사 열화 패턴 그룹화".encode("utf-8"), response.data)
        self.assertIn("패턴 분석 비활성".encode("utf-8"), response.data)
        self.assertIn(b'href="/inspections/ml-patterns"', response.data)
        self.assertIn(b'id="pattern-help-modal"', response.data)
        self.assertIn(b'data-pattern-help-open', response.data)
        self.assertIn("고변동·역전형".encode("utf-8"), response.data)
        self.assertIn("부식된 두께가 실제로 회복됐다는 의미가 아닙니다".encode("utf-8"), response.data)
        self.assertIn(b"pattern-help.js", response.data)
        script_response = self.client.get("/static/pattern-help.js")
        try:
            self.assertEqual(script_response.status_code, 200)
        finally:
            script_response.close()
        self.assertEqual(DB_PATH.read_bytes(), before)

    def test_exchanger_media_and_material_quick_management(self):
        with session_scope(DB_PATH) as session:
            before = session.scalar(select(HeatExchangerSpecRevision).where(
                HeatExchangerSpecRevision.heat_exchanger_id == 1,
                HeatExchangerSpecRevision.is_current.is_(True),
            ).order_by(HeatExchangerSpecRevision.revision_no.desc()))
            before_revision = before.revision_no if before else 0
            before_pressure = before.design_pressure if before else None

        response = self.client.post("/exchangers/1/media-material", data={
            "shell_fluid": "냉각수",
            "tube_fluid": "EDC",
            "tube_material": "Carbon Steel",
        }, follow_redirects=True)
        self.assertEqual(response.status_code, 200)
        self.assertIn("유체·재질 정보가 저장되었습니다.".encode("utf-8"), response.data)
        self.assertIn("냉각수".encode("utf-8"), response.data)
        self.assertIn(b"Carbon Steel", response.data)
        with session_scope(DB_PATH) as session:
            current = session.scalar(select(HeatExchangerSpecRevision).where(
                HeatExchangerSpecRevision.heat_exchanger_id == 1,
                HeatExchangerSpecRevision.is_current.is_(True),
            ).order_by(HeatExchangerSpecRevision.revision_no.desc()))
            self.assertEqual(current.revision_no, before_revision + 1)
            self.assertEqual(current.shell_fluid, "냉각수")
            self.assertEqual(current.tube_fluid, "EDC")
            self.assertEqual(current.tube_material, "Carbon Steel")
            self.assertEqual(current.design_pressure, before_pressure)

    def test_register_never_inspected_exchanger(self):
        new_page = self.client.get("/exchangers/new")
        self.assertEqual(new_page.status_code, 200)
        self.assertIn("유체·재질을 포함한 최초 Rev.1 등록".encode("utf-8"), new_page.data)
        self.assertIn(b'name="shell_fluid"', new_page.data)
        self.assertIn(b'name="tube_fluid"', new_page.data)
        self.assertIn(b'name="tube_material"', new_page.data)
        self.assertNotIn(b'name="manufacturing_date"', new_page.data)
        self.assertIn(b'name="equipment_tag"', new_page.data)
        self.assertNotIn(b'name="bundle_installed_at"', new_page.data)
        self.assertIn(b'name="head_material"', new_page.data)
        self.assertIn(b'name="head_fluid"', new_page.data)
        self.assertIn(b'name="tube_outer_diameter_mm"', new_page.data)
        self.assertIn("DB등록튜브수량".encode("utf-8"), new_page.data)
        self.assertIn(b'name="allowable_plugging_percent"', new_page.data)
        self.assertIn(b'type="range" name="allowable_plugging_percent"', new_page.data)
        self.assertIn(b'min="1" max="20" step="1" value="20"', new_page.data)
        response = self.client.post("/exchangers/new", data=self.mandatory_exchanger_form(
            equipment_code="E-TEST-V21", equipment_tag="US-VC-E-TEST-V21",
            exchanger_code="E-TEST-V21", equipment_name="Test exchanger",
            iris_applicability="OPTIONAL", expected_tube_count="20",
            process_criticality="CRITICAL", manufacturer="Maker A", model_no="HX-100",
            tube_material="Cu-Ni", design_pressure="10.5",
            inspection_cycle_months="24",
        ), follow_redirects=True)
        self.assertEqual(response.status_code, 200)
        self.assertIn("NEVER_INSPECTED".encode(), response.data)
        self.assertIn("E-TEST-V21".encode(), response.data)
        with session_scope(DB_PATH) as session:
            row = session.execute(
                select(HeatExchanger, Equipment)
                .join(Equipment, Equipment.equipment_id == HeatExchanger.equipment_id)
                .where(Equipment.equipment_code == "E-TEST-V21")
            ).one()
            exchanger_id = row.HeatExchanger.heat_exchanger_id
            self.assertEqual(row.Equipment.manufacturer, "Maker A")
            self.assertEqual(row.Equipment.equipment_tag, "US-VC-E-TEST-V21")
            self.assertEqual(
                row.Equipment.manufacturing_date.strftime("%Y-%m-%d"),
                "2005-11-28",
            )
            self.assertEqual(
                row.Equipment.commissioning_date.strftime("%Y-%m-%d"),
                "2005-12-01",
            )
            self.assertTrue(row.Equipment.is_critical)
            spec = session.scalar(select(HeatExchangerSpecRevision).where(
                HeatExchangerSpecRevision.heat_exchanger_id == exchanger_id,
                HeatExchangerSpecRevision.is_current.is_(True),
            ))
            criteria = session.scalar(select(HeatExchangerManagementCriteria).where(
                HeatExchangerManagementCriteria.heat_exchanger_id == exchanger_id
            ))
            self.assertEqual(spec.design_pressure, 10.5)
            self.assertEqual(spec.revision_no, 1)
            self.assertEqual(spec.shell_fluid, "Cooling Water")
            self.assertEqual(spec.tube_fluid, "Process EDC")
            self.assertEqual(spec.tube_material, "Cu-Ni")
            shell_side = session.scalar(select(HeatExchangerSideSpec).where(
                HeatExchangerSideSpec.spec_revision_id == spec.spec_revision_id,
                HeatExchangerSideSpec.side_code == "SHELL",
            ))
            head_side = session.scalar(select(HeatExchangerSideSpec).where(
                HeatExchangerSideSpec.spec_revision_id == spec.spec_revision_id,
                HeatExchangerSideSpec.side_code == "HEAD",
            ))
            self.assertEqual(shell_side.gasket_material, "Spiral wound")
            self.assertEqual(shell_side.concentration_wt_percent, 0.0)
            self.assertEqual(shell_side.material_state, "LIQUID")
            self.assertEqual(head_side.material, "SA-516-70")
            self.assertEqual(head_side.fluid, "Cooling Water")
            self.assertEqual(criteria.inspection_cycle_months, 24)
            self.assertEqual(criteria.allowable_plugging_percent, 20.0)

        critical_list = self.client.get("/exchangers?criticality=CRITICAL")
        self.assertEqual(critical_list.status_code, 200)
        self.assertIn(b"E-TEST-V21", critical_list.data)
        self.assertIn("공정 중요도".encode("utf-8"), critical_list.data)

        edit_response = self.client.post(
            f"/exchangers/{exchanger_id}/edit",
            data=self.mandatory_exchanger_form(
                equipment_code="E-TEST-V21", equipment_tag="US-VC-E-TEST-V21",
                exchanger_code="E-TEST-V21", equipment_name="Updated exchanger",
                manufacturer="Maker B", iris_applicability="OPTIONAL",
                install_date="2005-07-01",
                inspection_lifecycle_status="NEVER_INSPECTED",
                design_pressure="12.5", tube_material="Cu-Ni",
                warning_wall_loss_percent="25", critical_wall_loss_percent="45",
            ),
            follow_redirects=True,
        )
        self.assertEqual(edit_response.status_code, 200)
        self.assertIn("Updated exchanger".encode(), edit_response.data)
        with session_scope(DB_PATH) as session:
            equipment = session.scalar(select(Equipment).where(Equipment.equipment_code == "E-TEST-V21"))
            revision_count = session.scalar(select(func.count()).select_from(HeatExchangerSpecRevision).where(
                HeatExchangerSpecRevision.heat_exchanger_id == exchanger_id
            ))
            current = session.scalar(select(HeatExchangerSpecRevision).where(
                HeatExchangerSpecRevision.heat_exchanger_id == exchanger_id,
                HeatExchangerSpecRevision.is_current.is_(True),
            ))
            self.assertEqual(equipment.manufacturer, "Maker B")
            self.assertEqual(
                equipment.manufacturing_date.strftime("%Y-%m-%d"), "2005-06-28"
            )
            self.assertEqual(
                equipment.commissioning_date.strftime("%Y-%m-%d"), "2005-07-01"
            )
            self.assertEqual(revision_count, 2)
            self.assertEqual(current.revision_no, 2)
            self.assertEqual(current.design_pressure, 12.5)

    def test_exchanger_master_defined_fields_are_mandatory(self):
        page = self.client.get("/exchangers/new").get_data(as_text=True)
        for field_name in (
            "equipment_code", "equipment_tag", "exchanger_code", "equipment_name",
            "service_name", "exchanger_type_code", "status", "install_date",
            "shell_material", "head_material", "tube_material", "tube_count",
            "tube_outer_diameter_mm", "tube_nominal_thickness_mm", "tube_length_mm",
            "tube_pitch_mm", "pass_count", "shell_fluid", "head_fluid", "tube_fluid",
        ):
            self.assertRegex(page, rf'name="{field_name}"[^>]*required')

        incomplete = self.mandatory_exchanger_form(
            equipment_code="E-MISSING-MASTER",
            equipment_tag="US-VC-E-MISSING-MASTER",
            exchanger_code="E-MISSING-MASTER",
            head_material="",
        )
        response = self.client.post(
            "/exchangers/new", data=incomplete, follow_redirects=True
        )
        self.assertIn("마스터 필수 항목을 입력해 주세요: Head 재질".encode("utf-8"), response.data)
        with session_scope(DB_PATH) as session:
            self.assertIsNone(session.scalar(select(Equipment).where(
                Equipment.equipment_code == "E-MISSING-MASTER"
            )))

    def test_inspection_plan_requires_in_service_tubes_and_admin_approval(self):
        with session_scope(DB_PATH) as session:
            admin_id = session.scalar(select(AppUser.user_id).where(
                AppUser.is_admin.is_(True),
                AppUser.is_active.is_(True),
            ))
            tube = session.scalar(
                select(HeatExchangerTube).join(TubeBundle).where(
                    TubeBundle.heat_exchanger_id == 1,
                    HeatExchangerTube.current_status.in_(["ACTIVE", "IN_SERVICE"]),
                ).limit(1)
            )
            tube_id = tube.tube_id

        selection_page = self.client.get("/inspection-plans?heat_exchanger_id=1")
        self.assertEqual(selection_page.status_code, 200)
        self.assertIn("사용 중 Tube 선택".encode("utf-8"), selection_page.data)
        self.assertIn(f'value="{tube_id}"'.encode(), selection_page.data)

        plan_id = services.save_inspection_plan({
            "heat_exchanger_id": "1",
            "planned_date": "2027-06-01",
            "inspection_method": "IRIS",
            "scope_code": "ALL",
            "priority": "NORMAL",
            "tube_ids": [str(tube_id)],
        }, requested_by_user_id=admin_id)
        with session_scope(DB_PATH) as session:
            plan = session.get(InspectionPlan, plan_id)
            target = session.scalar(select(InspectionPlanTubeTarget).where(
                InspectionPlanTubeTarget.inspection_plan_id == plan_id
            ))
            self.assertEqual(plan.requested_by_user_id, admin_id)
            self.assertEqual(target.tube_status_at_selection, "IN_SERVICE")
            target_id = target.inspection_plan_tube_target_id

        with self.assertRaisesRegex(ValueError, "관리자만"):
            services.update_inspection_plan_status(
                plan_id, "APPROVED", actor_user_id=admin_id, actor_is_admin=False
            )
        with session_scope(DB_PATH) as session:
            session.get(HeatExchangerTube, tube_id).current_status = "PLUGGED"
        with self.assertRaisesRegex(ValueError, "현재 사용 중이 아닌"):
            services.update_inspection_plan_status(
                plan_id, "APPROVED", actor_user_id=admin_id, actor_is_admin=True
            )
        with session_scope(DB_PATH) as session:
            session.get(HeatExchangerTube, tube_id).current_status = "IN_SERVICE"
        self.assertTrue(services.update_inspection_plan_status(
            plan_id, "APPROVED", actor_user_id=admin_id, actor_is_admin=True,
            reason="검사 범위 확인",
        ))
        self.assertTrue(services.update_inspection_plan_status(
            plan_id, "IN_PROGRESS", actor_user_id=admin_id, actor_is_admin=True
        ))
        with self.assertRaisesRegex(ValueError, "미검사 대상 Tube"):
            services.update_inspection_plan_status(
                plan_id, "COMPLETED", actor_user_id=admin_id, actor_is_admin=True
            )
        self.assertTrue(services.update_inspection_plan_tube_target(
            plan_id, target_id, "INSPECTED", actor_user_id=admin_id
        ))
        self.assertTrue(services.update_inspection_plan_status(
            plan_id, "COMPLETED", actor_user_id=admin_id, actor_is_admin=True
        ))
        with session_scope(DB_PATH) as session:
            plan = session.get(InspectionPlan, plan_id)
            self.assertEqual(plan.status, "COMPLETED")
            self.assertEqual(plan.approved_by_user_id, admin_id)
            self.assertIsNotNone(plan.target_locked_at)
            self.assertGreaterEqual(session.scalar(select(func.count()).select_from(
                InspectionPlanEvent
            ).where(InspectionPlanEvent.inspection_plan_id == plan_id)), 4)

        with session_scope(DB_PATH) as session:
            tube = session.get(HeatExchangerTube, tube_id)
            original_status = tube.current_status
            tube.current_status = "PLUGGED"
        try:
            with self.assertRaisesRegex(ValueError, "사용 중"):
                services.save_inspection_plan({
                    "heat_exchanger_id": "1", "planned_date": "2027-07-01",
                    "inspection_method": "IRIS", "scope_code": "ALL",
                    "priority": "NORMAL", "tube_ids": [str(tube_id)],
                }, requested_by_user_id=admin_id)
        finally:
            with session_scope(DB_PATH) as session:
                session.get(HeatExchangerTube, tube_id).current_status = original_status

    def test_historical_inspection_results_backfill_complete_plan_workflow(self):
        with session_scope(DB_PATH) as session:
            applicant = session.scalar(select(AppUser).where(
                AppUser.username == "history.applicant"
            ))
            manager = session.scalar(select(AppUser).where(
                AppUser.username == "history.manager"
            ))
            self.assertIsNotNone(applicant)
            self.assertIsNotNone(manager)
            self.assertFalse(applicant.is_active)
            self.assertFalse(applicant.is_admin)
            self.assertFalse(manager.is_active)
            self.assertTrue(manager.is_admin)

            plans = session.scalars(select(InspectionPlan).where(
                InspectionPlan.completed_inspection_id.is_not(None)
            )).all()
            result_count = session.scalar(
                select(func.count()).select_from(IrisTubeResult)
            ) or 0
            target_count = session.scalar(select(func.count()).select_from(
                InspectionPlanTubeTarget
            ).where(
                InspectionPlanTubeTarget.source_iris_tube_result_id.is_not(None),
                InspectionPlanTubeTarget.target_status == "INSPECTED",
            )) or 0
            self.assertEqual(target_count, result_count)
            self.assertTrue(plans)
            for plan in plans:
                self.assertEqual(plan.status, "COMPLETED")
                self.assertEqual(plan.requested_by_user_id, applicant.user_id)
                self.assertEqual(plan.approved_by_user_id, manager.user_id)
                self.assertIsNotNone(plan.application_due_date)
                self.assertIsNotNone(plan.target_locked_at)
                self.assertIsNotNone(plan.inspection_started_at)
                self.assertIsNotNone(plan.inspection_completed_at)
                event_count = session.scalar(
                    select(func.count()).select_from(InspectionPlanEvent).where(
                        InspectionPlanEvent.inspection_plan_id == plan.inspection_plan_id
                    )
                ) or 0
                self.assertGreaterEqual(event_count, 4)

    def test_priority_restored_workflows(self):
        for path in ("/inspections", "/inspections/tubes", "/inspection-plans", "/maintenance",
                     "/exchangers/1/report", "/exchangers/1/tubes/manage",
                     "/api/v1/exchangers", "/api/v1/exchangers/1/history"):
            self.assertEqual(self.client.get(path).status_code, 200, path)

        with session_scope(DB_PATH) as session:
            scheduled_tube_ids = list(session.scalars(
                select(HeatExchangerTube.tube_id)
                .join(TubeBundle)
                .where(
                    TubeBundle.heat_exchanger_id == 1,
                    HeatExchangerTube.current_status.in_(["ACTIVE", "IN_SERVICE"]),
                ).limit(2)
            ))
        plan_response = self.client.post("/inspection-plans", data={
            "heat_exchanger_id": "1", "planned_date": "2027-05-01",
            "inspection_method": "IRIS", "scope_code": "ALL", "priority": "HIGH",
            "tube_ids": [str(value) for value in scheduled_tube_ids],
        }, follow_redirects=True)
        self.assertIn("검사계획이 등록되었습니다.".encode("utf-8"), plan_response.data)
        with session_scope(DB_PATH) as session:
            plan = session.scalar(select(InspectionPlan).order_by(InspectionPlan.inspection_plan_id.desc()))
            plan_id = plan.inspection_plan_id
            self.assertEqual(session.scalar(select(func.count()).select_from(
                InspectionPlanTubeTarget
            ).where(InspectionPlanTubeTarget.inspection_plan_id == plan_id)), 2)
        status_response = self.client.post(f"/inspection-plans/{plan_id}/status",
                                           data={"status": "APPROVED"}, follow_redirects=True)
        self.assertIn(b"APPROVED", status_response.data)

        with session_scope(DB_PATH) as session:
            tube_ids = list(session.scalars(select(HeatExchangerTube.tube_id)
                .join(TubeBundle).where(TubeBundle.heat_exchanger_id == 1).limit(2)))
        batch = self.client.post("/exchangers/1/tubes/batch-status", data={
            "tube_ids": [str(value) for value in tube_ids], "status": "PLUGGED",
            "reason": "integration test",
        }, follow_redirects=True)
        self.assertIn("튜브 2개의 상태".encode("utf-8"), batch.data)
        position = self.client.post(f"/tubes/{tube_ids[0]}/position", data={
            "position_x": "10.5", "position_y": "20.5", "position_z": "1",
            "position_locked": "1", "reason": "drawing verified",
        }, follow_redirects=True)
        self.assertIn("좌표와 잠금 상태".encode("utf-8"), position.data)

        work = self.client.post("/maintenance/work-orders/new", data={
            "heat_exchanger_id": "1", "work_order_code": "WO-RESTORE-001",
            "title": "열교환기 전체 교체", "work_type": "REPLACEMENT",
        }, follow_redirects=True)
        self.assertIn(b"WO-RESTORE-001", work.data)
        with session_scope(DB_PATH) as session:
            order = session.scalar(select(WorkOrder).where(WorkOrder.work_order_code == "WO-RESTORE-001"))
            order_id = order.work_order_id
            foreign_tube_id = session.scalar(
                select(HeatExchangerTube.tube_id)
                .join(TubeBundle)
                .where(TubeBundle.heat_exchanger_id != 1)
                .limit(1)
            )
            record_count_before = session.scalar(
                select(func.count()).select_from(MaintenanceRecord)
            )
        order_page = self.client.get(f"/maintenance/work-orders/{order_id}")
        self.assertIn("PLUGGED 상태 제외".encode("utf-8"), order_page.data)
        self.assertIn(b"maintenance-record.js", order_page.data)
        self.assertIn(f'value="{tube_ids[0]}"'.encode(), order_page.data)

        negative_id = self.client.post(
            f"/maintenance/work-orders/{order_id}/records",
            data={"tube_id": "-1", "action_type": "REPLACEMENT"},
            follow_redirects=True,
        )
        self.assertIn("Tube ID는 1 이상의 양수".encode("utf-8"), negative_id.data)
        foreign_id = self.client.post(
            f"/maintenance/work-orders/{order_id}/records",
            data={"tube_id": str(foreign_tube_id), "action_type": "REPLACEMENT"},
            follow_redirects=True,
        )
        self.assertIn("이 작업지시의 열교환기에 등록된 튜브가 아닙니다".encode("utf-8"), foreign_id.data)
        with session_scope(DB_PATH) as session:
            self.assertEqual(
                session.scalar(select(func.count()).select_from(MaintenanceRecord)),
                record_count_before,
            )
        record = self.client.post(f"/maintenance/work-orders/{order_id}/records", data={
            "tube_id": str(tube_ids[0]), "action_type": "MAINTENANCE",
            "new_status": "IN_SERVICE",
            "details": "maintenance completed",
        }, follow_redirects=True)
        self.assertIn("maintenance completed".encode(), record.data)
        rejected = self.client.post(f"/maintenance/work-orders/{order_id}/records", data={
            "tube_id": str(tube_ids[0]), "action_type": "REPLACEMENT",
        }, follow_redirects=True)
        self.assertIn("튜브를 선택할 수 없습니다".encode(), rejected.data)
        with session_scope(DB_PATH) as session:
            self.assertEqual(session.scalar(select(func.count()).select_from(MaintenanceRecord)),
                             record_count_before + 1)
        replacement = self.client.post(f"/maintenance/work-orders/{order_id}/records", data={
            "action_type": "REPLACEMENT", "details": "exchanger replacement completed",
        }, follow_redirects=True)
        self.assertIn(b"exchanger replacement completed", replacement.data)
        with session_scope(DB_PATH) as session:
            replacement_record = session.scalar(select(MaintenanceRecord).where(
                MaintenanceRecord.work_order_id == order_id,
                MaintenanceRecord.action_type == "REPLACEMENT"))
            self.assertIsNotNone(replacement_record)
            self.assertIsNone(replacement_record.tube_id)
            self.assertIsNone(replacement_record.new_status)
        with session_scope(DB_PATH) as session:
            self.assertGreaterEqual(session.scalar(select(func.count()).select_from(TubePositionEvent)), 1)
            self.assertGreaterEqual(session.scalar(select(func.count()).select_from(MaintenanceRecord)), 1)

        self.assertEqual(self.client.post("/api/v1/exchangers", json={}).status_code, 503)
        with patch("iris2.app.API_WRITE_TOKEN", "test-token"):
            api_create = self.client.post("/api/v1/exchangers", headers={
                "X-IRIS-API-Token": "test-token"
            }, json={"unit_id": 1, "equipment_code": "E-API-001",
                     "equipment_name": "API exchanger"})
            self.assertEqual(api_create.status_code, 201)
            api_id = api_create.get_json()["heat_exchanger_id"]
            api_update = self.client.put(f"/api/v1/exchangers/{api_id}", headers={
                "X-IRIS-API-Token": "test-token"
            }, json={"unit_id": 1, "equipment_code": "E-API-001",
                     "equipment_name": "API exchanger updated"})
            self.assertEqual(api_update.status_code, 200)

    def test_plant_and_production_line_master_can_supply_new_exchanger(self):
        page = self.client.get("/master-data/plants-lines")
        self.assertEqual(page.status_code, 200)
        self.assertIn("공장·생산라인 코드 관리".encode("utf-8"), page.data)
        self.assertIn("VCM1".encode(), page.data)

        plant_response = self.client.post("/master-data/plants", data={
            "plant_code": "P2",
            "plant_name": "제2공장",
            "location": "여수",
        }, follow_redirects=True)
        self.assertEqual(plant_response.status_code, 200)
        self.assertIn("공장 코드가 등록되었습니다.".encode("utf-8"), plant_response.data)
        self.assertIn("P1부터 Pn까지".encode("utf-8"), plant_response.data)

        with session_scope(DB_PATH) as session:
            plant_id = session.scalar(select(Plant.plant_id).where(Plant.plant_code == "P2"))
        line_response = self.client.post("/master-data/production-lines", data={
            "plant_id": str(plant_id),
            "unit_code": "LINE-A",
            "unit_name": "A 생산라인",
        }, follow_redirects=True)
        self.assertEqual(line_response.status_code, 200)
        self.assertIn("생산라인 코드가 등록되었습니다.".encode("utf-8"), line_response.data)
        import_page = self.client.get("/imports").get_data(as_text=True)
        self.assertIn('value="P2"', import_page)
        self.assertIn("생산공장을 선택해 주세요", import_page)
        self.assertIn("생산라인을 선택해 주세요", import_page)
        self.assertNotIn('value="P1" selected', import_page)
        self.assertNotIn('value="P2" selected', import_page)

        with session_scope(DB_PATH) as session:
            unit_id = session.scalar(select(ProcessUnit.unit_id).where(
                ProcessUnit.plant_id == plant_id,
                ProcessUnit.unit_code == "LINE-A",
            ))
        create_response = self.client.post(
            "/exchangers/new",
            data=self.mandatory_exchanger_form(
                unit_id=str(unit_id), plant_code="P2", unit_code="LINE-A",
                equipment_code="E-LINE-A-001", equipment_tag="US-LINE-A-E-001",
                exchanger_code="E-LINE-A-001", equipment_name="신규 라인 열교환기",
                tube_count="16",
            ),
            follow_redirects=True,
        )
        self.assertEqual(create_response.status_code, 200)
        self.assertIn("E-LINE-A-001".encode(), create_response.data)
        with session_scope(DB_PATH) as session:
            equipment = session.scalar(select(Equipment).where(
                Equipment.equipment_code == "E-LINE-A-001"
            ))
            self.assertEqual(equipment.unit_id, unit_id)
            line = session.get(ProcessUnit, equipment.unit_id)
            self.assertEqual(line.unit_code, "LINE-A")
            self.assertEqual(line.plant_id, plant_id)

    def test_json_load_requires_explicit_plant_selection(self):
        response = self.client.post("/imports", data={
            "json_file": (BytesIO(b"{}"), "no-plant.json"),
        }, content_type="multipart/form-data", follow_redirects=True)
        self.assertEqual(response.status_code, 200)
        self.assertIn("생산공장 코드는 필수입니다.".encode("utf-8"), response.data)

        invalid_plant = self.client.post("/master-data/plants", data={
            "plant_code": "FACTORY3",
            "plant_name": "잘못된 코드",
        }, follow_redirects=True)
        self.assertIn("P1, P2, P3".encode("utf-8"), invalid_plant.data)
        with session_scope(DB_PATH) as session:
            self.assertIsNone(session.scalar(select(Plant).where(Plant.plant_code == "FACTORY3")))

    def test_exchanger_rejects_unregistered_plant_line_codes(self):
        response = self.client.post("/exchangers/new", data={
            "plant_code": "TYPO-PLANT",
            "unit_code": "TYPO-LINE",
            "equipment_code": "E-TYPO-001",
            "equipment_name": "오입력 방지 확인",
        }, follow_redirects=True)
        self.assertEqual(response.status_code, 200)
        self.assertIn("기준정보에서 먼저 등록".encode("utf-8"), response.data)
        with session_scope(DB_PATH) as session:
            self.assertIsNone(session.scalar(select(Plant).where(Plant.plant_code == "TYPO-PLANT")))

    def test_unknown_records_return_404(self):
        self.assertEqual(self.client.get("/exchangers/999999").status_code, 404)
        self.assertEqual(self.client.get("/inspections/999999").status_code, 404)
        self.assertEqual(self.client.get("/tubes/999999").status_code, 404)

    def test_load_result_modal_renders_all_load_metrics(self):
        outcome = {
            "status": "COMPLETED_WITH_WARNINGS",
            "source_file": "2022VCM1_IRIS.json",
            "total_exchanger_count": 12,
            "new_exchanger_count": 4,
            "existing_exchanger_count": 8,
            "missing_master_exchanger_count": 4,
            "missing_equipment_messages": [
                "미존재 열교환기: VCM1+E-404 - Master DB에 없어 적재 제외"
            ],
            "report_count": 20,
            "loaded_report_count": 20,
            "tube_result_count": 350,
            "error_count": 0,
            "warning_count": 3,
            "degradation_assessment_count": 120,
        }
        with self.app.test_request_context("/imports"):
            html = render_template("imports.html", outcome=outcome)
        self.assertIn('id="load-result-modal"', html)
        self.assertIn("DB 적재 완료", html)
        self.assertIn("2022VCM1_IRIS.json", html)
        self.assertIn("이력·감육 갱신", html)
        self.assertIn("350", html)
        self.assertIn("Master 미존재", html)
        self.assertIn("기존 열교환기", html)
        self.assertIn("미존재 열교환기: VCM1+E-404", html)
        self.assertIn(">4</strong>", html)
        self.assertIn(">8</strong>", html)

    def test_json_import_requires_registered_production_line(self):
        page = self.client.get("/imports")
        self.assertEqual(page.status_code, 200)
        self.assertIn(b'name="plant_code"', page.data)
        self.assertIn(b'name="unit_code"', page.data)
        self.assertIn(b'data-plant="P1"', page.data)
        self.assertIn(b'import-location.js', page.data)
        with self.assertRaisesRegex(ValueError, "등록된 생산라인"):
            services.load_json(
                APP_DIR / "data" / "2020VCM1_IRIS2_all_tube_results.json",
                dry_run=True,
                plant_code="P1",
                unit_code="NOT_REGISTERED",
            )
        with patch("iris2.services.import_file", return_value={"unit_code": "VCM1"}) as importer:
            services.load_json(
                APP_DIR / "data" / "unused.json",
                dry_run=True,
                plant_code="P1",
                unit_code="VCM1",
            )
        self.assertEqual(importer.call_args.kwargs["plant_code"], "P1")
        self.assertEqual(importer.call_args.kwargs["unit_code"], "VCM1")

    def test_dashboard_pdf_report_link_opens_inline(self):
        dashboard = self.client.get("/").get_data(as_text=True)
        self.assertIn("Report View", dashboard)
        self.assertIn('/pdf-reports/2020/VCM1', dashboard)
        response = self.client.get("/pdf-reports/2020/VCM1")
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.mimetype, "application/pdf")
        self.assertIn("inline", response.headers.get("Content-Disposition", ""))
        response.close()
        self.assertEqual(self.client.get("/pdf-reports/2024/VCM1").status_code, 404)
        self.assertEqual(self.client.get("/pdf-reports/2020/not-a-vcm").status_code, 404)

    def test_exchanger_management_links_inspection_sheet_pdf(self):
        page = self.client.get("/exchangers").get_data(as_text=True)
        self.assertIn("검사 Sheet", page)
        self.assertRegex(page, r'/inspections/\d+/pdf')
        pdf_link_tag = page.split('class="button report-view"', 1)[1].split(">", 1)[0]
        self.assertIn('target="_blank"', pdf_link_tag)
        self.assertIsNotNone(services.find_pdf_report(2020, "PE"))

        response = self.client.get("/inspections/1/pdf")
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.mimetype, "application/pdf")
        self.assertIn("inline", response.headers.get("Content-Disposition", ""))
        response.close()
        self.assertEqual(self.client.get("/inspections/999999/pdf").status_code, 404)

    def test_excel_export_contains_exchangers_tubes_and_history(self):
        with session_scope(DB_PATH) as session:
            expected_exchangers = session.scalar(select(func.count()).select_from(HeatExchanger))
            expected_tubes = session.scalar(select(func.count()).select_from(HeatExchangerTube))
            expected_results = session.scalar(select(func.count()).select_from(IrisTubeResult))
        dashboard = self.client.get("/").get_data(as_text=True)
        topbar = dashboard[dashboard.index('<header class="topbar">'):dashboard.index("</header>")]
        self.assertNotIn("Excel Export", topbar)
        self.assertNotIn("JSON 신규 적재", topbar)
        self.assertIn("Excel Export", dashboard[dashboard.index("<nav>"):dashboard.index("</nav>")])
        self.assertIn("sidebar-toggle", dashboard)
        self.assertIn("sidebar.js", dashboard)
        sidebar_script = self.client.get("/static/sidebar.js")
        self.assertEqual(sidebar_script.status_code, 200)
        self.assertIn(b"sidebar-collapsed", sidebar_script.data)
        sidebar_script.close()
        response = self.client.get("/exports/heat-exchangers-tubes.xlsx")
        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            response.mimetype,
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        self.assertIn("attachment", response.headers.get("Content-Disposition", ""))
        wb = load_workbook(BytesIO(response.data), read_only=False, data_only=False)
        self.assertEqual(
            wb.sheetnames,
            ["요약", "열교환기", "튜브", "검사보고서", "검사이력", "상태변경이력"],
        )
        self.assertEqual(wb["열교환기"].max_row - 4, expected_exchangers)
        self.assertEqual(wb["튜브"].max_row - 4, expected_tubes)
        self.assertEqual(wb["검사이력"].max_row - 4, expected_results)
        self.assertEqual(wb["요약"]["B5"].value, "=COUNTA('열교환기'!$A$5:$A$2000)")
        exchanger_headers = [cell.value for cell in wb["열교환기"][4]]
        self.assertIn("제작일자", exchanger_headers)
        self.assertIn("최초 가동일", exchanger_headers)
        self.assertIn("열교환기코드", exchanger_headers)
        self.assertEqual(wb["튜브"].freeze_panes, "A5")
        self.assertTrue(wb["튜브"].tables)
        wb.close()

    def test_exchanger_tube_map_2d_3d_data(self):
        page = self.client.get("/exchangers/1/tube-map")
        self.assertEqual(page.status_code, 200)
        self.assertIn(b"2D Map", page.data)
        self.assertIn(b"3D Map", page.data)
        self.assertIn("색상 기준".encode("utf-8"), page.data)
        self.assertIn("튜브 내면·외면 부식 펼침도".encode("utf-8"), page.data)
        self.assertIn("단일 튜브 조작".encode("utf-8"), page.data)
        self.assertIn(b'id="focus-zoom-in"', page.data)
        self.assertIn(b'id="focus-reset"', page.data)
        self.assertIn("전체 회전".encode("utf-8"), page.data)
        self.assertIn("튜브 이동".encode("utf-8"), page.data)
        self.assertIn("번호·Row/Column 표시 및 축 방향 자동 인출".encode("utf-8"), page.data)
        self.assertIn(b"TransformControls.js", page.data)
        response = self.client.get("/api/exchangers/1/tube-map")
        self.assertEqual(response.status_code, 200)
        payload = response.get_json()
        self.assertGreater(len(payload["tubes"]), 0)
        self.assertGreater(len(payload["campaigns"]), 0)
        category_total = sum(payload["summary"][key] for key in (
            "NOT_INSPECTED", "NORMAL", "WARNING", "CRITICAL", "PLUGGED"
        ))
        self.assertEqual(category_total, payload["summary"]["display_tube_count"])
        self.assertEqual(len(payload["tubes"]), payload["summary"]["expected_tube_count"])
        self.assertEqual(len(payload["tubes"]), payload["summary"]["display_tube_count"])
        self.assertEqual(
            sum(1 for tube in payload["tubes"] if not tube["managed"]),
            payload["summary"]["unmapped_expected_count"],
        )
        self.assertTrue(payload["layout"]["is_estimated"])
        self.assertIn("orientation", payload["layout"])
        self.assertIn(payload["layout"]["initial_rotation_degrees"], (0, 90, 180, 270))
        self.assertIn("tube_spec", payload)
        self.assertIn("length_mm", payload["tube_spec"])
        self.assertIn(payload["tube_spec"]["source"], (
            "MASTER", "INSPECTION_SNAPSHOT", "EXCHANGER_AVERAGE",
            "GLOBAL_AVERAGE", "UNREGISTERED",
        ))
        self.assertGreater(payload["tube_spec"]["length_mm"], 0)
        self.assertIn("source_label", payload["tube_spec"])
        self.assertIn("length_is_estimated", payload["tube_spec"])
        self.assertTrue(all("wall_loss_percent" in tube for tube in payload["tubes"]))
        self.assertTrue(all("position_locked" in tube for tube in payload["tubes"] if tube["managed"]))
        self.assertTrue(all("coordinate_quality" in tube for tube in payload["tubes"]))
        self.assertIn("data_quality", payload)
        self.assertIn("registered_xy_count", payload["data_quality"])
        self.assertIn("measured_axial_position_count", payload["data_quality"])
        self.assertIn("circumferential_angle_count", payload["data_quality"])
        self.assertIn("defect_length_count", payload["data_quality"])
        self.assertEqual(payload["thresholds"], {"warning": 20.0, "critical": 40.0})
        self.assertEqual(payload["colors"]["NORMAL"], "#14b8a6")
        self.assertEqual(payload["colors"]["CRITICAL"], "#ef4444")
        movable = next(tube for tube in payload["tubes"] if tube["managed"] and not tube["position_locked"])
        move = self.client.post(f"/api/tubes/{movable['tube_id']}/position", json={
            "position_x": 12.25, "position_y": 8.75, "position_z": 0,
            "position_locked": False, "reason": "2D drag test",
        })
        self.assertEqual(move.status_code, 200)
        with session_scope(DB_PATH) as session:
            moved = session.get(HeatExchangerTube, movable["tube_id"])
            self.assertEqual(moved.position_x, 12.25)
            self.assertEqual(moved.position_y, 8.75)
        self.assertEqual(self.client.get("/api/exchangers/999999/tube-map").status_code, 404)
        three = self.client.get("/assets/three.min.js")
        self.assertEqual(three.status_code, 200)
        self.assertEqual(three.mimetype, "application/javascript")
        three.close()
        transform = self.client.get("/assets/TransformControls.js")
        self.assertEqual(transform.status_code, 200)
        self.assertEqual(transform.mimetype, "application/javascript")
        transform.close()
        self.assertIn("3D 단일 튜브 분리".encode(), page.data)
        self.assertIn("실제 Tube 길이".encode(), page.data)
        self.assertIn("검사 위치·측정값".encode(), page.data)
        self.assertIn("튜브 내면·외면 부식 펼침도".encode(), page.data)
        self.assertIn(b'id="tube-unwrapped-inner"', page.data)
        self.assertIn(b'id="tube-unwrapped-outer"', page.data)
        self.assertIn(b'id="tube-unwrapped-unknown"', page.data)
        self.assertIn("위치 기준".encode(), page.data)
        self.assertIn(b"data-pdf-url-template", page.data)
        self.assertTrue(all("indications" in tube for tube in payload["tubes"]))
        self.assertTrue(all(
            tube["indication_count"] == len(tube["indications"])
            for tube in payload["tubes"]
        ))
        inspected = next(tube for tube in payload["tubes"] if tube["inspected"])
        self.assertGreater(inspected["indication_count"], 0)
        indication = inspected["indications"][0]
        self.assertIn(indication["surface_side"], ("INNER", "OUTER", "UNKNOWN"))
        self.assertIn(indication["position_quality"], (
            "MEASURED", "RESULT_REPRESENTATIVE", "UNKNOWN",
        ))
        self.assertIn("position_reference", indication)
        self.assertIn("inspection_id", indication)
        self.assertIn("source_page", indication)
        self.assertEqual(
            inspected["measured_indication_count"],
            sum(
                item["position_quality"] == "MEASURED"
                for item in inspected["indications"]
            ),
        )
        self.assertEqual(
            inspected["circumferential_indication_count"],
            sum(
                item["circumferential_angle_deg"] is not None
                for item in inspected["indications"]
            ),
        )

        script_response = self.client.get("/static/tube-map.js")
        script = script_response.get_data(as_text=True)
        self.assertIn("initial_rotation_degrees", script)
        self.assertIn("위치 미확정", script)
        self.assertIn("drawUnwrappedTube", script)
        self.assertIn("drawUnwrappedSurface", script)
        self.assertIn("tube-selection-badge", script)
        self.assertIn("pullDistance", script)
        self.assertIn("moved>6", script)
        self.assertIn("tubeIdentityLabel", script)
        self.assertIn("surface-data-warning", script)
        self.assertIn("equivalentWallLoss", script)
        self.assertIn("rotateFocusedTube", script)
        self.assertIn("zoomFocusedTube", script)
        self.assertIn("screenSpacePanning=true", script)
        script_response.close()

    def test_indication_surface_side_requires_explicit_source_value(self):
        explicit_inner = SimpleNamespace(
            defect_type="PIT", judgement=None,
            raw_json=json.dumps({"data": {"surface_side": "ID"}}),
        )
        explicit_outer = SimpleNamespace(
            defect_type="COR", judgement=None,
            raw_json=json.dumps({"corrosion_side": "external"}),
        )
        unknown = SimpleNamespace(
            defect_type="WLL", judgement=None,
            raw_json=json.dumps({"data": {"SurfacePercent": 0}}),
        )
        self.assertEqual(services._indication_surface_side(explicit_inner), "INNER")
        self.assertEqual(services._indication_surface_side(explicit_outer), "OUTER")
        self.assertEqual(services._indication_surface_side(unknown), "UNKNOWN")

    def test_tube_map_orientation_rotation_rules(self):
        self.assertEqual(services._tube_map_initial_rotation_degrees("Front View"), 0)
        self.assertEqual(services._tube_map_initial_rotation_degrees("Top(North) View"), 0)
        self.assertEqual(services._tube_map_initial_rotation_degrees("Front(East) View"), 90)
        self.assertEqual(services._tube_map_initial_rotation_degrees("Rear(South) View"), 180)
        self.assertEqual(services._tube_map_initial_rotation_degrees("Top West"), 270)
        self.assertEqual(services._tube_map_initial_rotation_degrees(None), 0)

    def test_tube_map_uses_approved_layout_without_changing_tube_master_xy(self):
        with session_scope(DB_PATH) as session:
            bundle = session.scalar(
                select(TubeBundle)
                .where(TubeBundle.heat_exchanger_id == 1)
                .order_by(TubeBundle.bundle_id)
            )
            tube = session.scalar(
                select(HeatExchangerTube)
                .where(HeatExchangerTube.bundle_id == bundle.bundle_id)
                .order_by(HeatExchangerTube.tube_id)
            )
            original_xy = (tube.position_x, tube.position_y)
            revision = TubeLayoutRevision(
                bundle_id=bundle.bundle_id,
                source_path="test/approved-layout.mdb",
                source_section_id=1,
                source_section_name="Front View",
                source_coordinate_unit="GRID",
                rotation_degrees=0,
                physical_tube_count=1,
                layout_point_count=2,
                layout_count_basis="TUBE_END",
                coverage_status="FULL",
                unresolved_count=0,
                revision_status="APPROVED",
                approved_by="test",
                approved_at=datetime.now(),
            )
            session.add(revision)
            session.flush()
            revision_id = revision.layout_revision_id
            tube_id = tube.tube_id
            session.add_all([
                TubeLayoutPoint(
                    layout_revision_id=revision.layout_revision_id,
                    tube_id=tube.tube_id,
                    source_tube_id=tube.tube_no,
                    source_row=tube.row_no,
                    source_column=tube.column_no,
                    source_x=101,
                    source_y=202,
                    mapping_status="EXISTING_TUBE_MATCH",
                    verification_status="VERIFIED",
                ),
                TubeLayoutPoint(
                    layout_revision_id=revision.layout_revision_id,
                    source_tube_id="layout-only",
                    source_row=999,
                    source_column=999,
                    source_x=303,
                    source_y=404,
                    mapping_status="NEW_TUBE_CANDIDATE",
                    verification_status="VERIFIED",
                ),
            ])

        payload = self.client.get("/api/exchangers/1/tube-map").get_json()
        self.assertTrue(payload["layout"]["approved_layout_applied"])
        self.assertEqual(payload["layout"]["shape"], "SOURCE_LAYOUT_GRID")
        self.assertEqual(payload["layout"]["initial_rotation_degrees"], 0)
        self.assertEqual(payload["summary"]["display_tube_count"], 2)
        self.assertEqual(payload["summary"]["display_position_count"], 2)
        self.assertEqual(payload["summary"]["expected_tube_count"], 1)
        self.assertEqual(payload["summary"]["unmapped_expected_count"], 0)
        self.assertEqual(payload["summary"]["unlinked_layout_count"], 1)
        self.assertEqual(payload["layout"]["physical_tube_count"], 1)
        self.assertEqual(payload["layout"]["layout_point_count"], 2)
        self.assertEqual(payload["layout"]["count_basis"], "TUBE_END")
        self.assertEqual(payload["layout"]["coverage_status"], "FULL")
        linked = next(item for item in payload["tubes"] if item["tube_id"] == tube_id)
        self.assertTrue(linked["managed"])
        self.assertEqual((linked["x"], linked["y"]), (101.0, 202.0))
        self.assertEqual(linked["coordinate_source"], "APPROVED_LAYOUT")
        self.assertTrue(linked["position_locked"])
        with session_scope(DB_PATH) as session:
            unchanged = session.get(HeatExchangerTube, tube_id)
            self.assertEqual((unchanged.position_x, unchanged.position_y), original_xy)
            for point in session.scalars(
                select(TubeLayoutPoint).where(
                    TubeLayoutPoint.layout_revision_id == revision_id
                )
            ).all():
                session.delete(point)
            session.delete(session.get(TubeLayoutRevision, revision_id))

    def test_wall_loss_color_settings_can_update_and_reset(self):
        page = self.client.get("/settings/wall-loss-colors")
        self.assertEqual(page.status_code, 200)
        self.assertIn("Wall Loss 색상표".encode(), page.data)
        self.assertIn(b"#14b8a6", page.data)

        custom = {
            "NORMAL": "#00aa88",
            "WARNING": "#ffaa00",
            "CRITICAL": "#cc2233",
            "NOT_INSPECTED": "#778899",
            "PLUGGED": "#222222",
        }
        response = self.client.post("/settings/wall-loss-colors", data=custom, follow_redirects=True)
        self.assertEqual(response.status_code, 200)
        self.assertIn("변경 5건".encode(), response.data)
        payload = self.client.get("/api/exchangers/1/tube-map").get_json()
        self.assertEqual(payload["colors"], {key: value.lower() for key, value in custom.items()})
        with session_scope(DB_PATH) as session:
            self.assertEqual(session.get(WallLossColorSetting, "NORMAL").current_color, "#00aa88")

        reset = self.client.post("/settings/wall-loss-colors/reset", follow_redirects=True)
        self.assertEqual(reset.status_code, 200)
        self.assertIn("복원 5건".encode(), reset.data)
        payload = self.client.get("/api/exchangers/1/tube-map").get_json()
        self.assertEqual(payload["colors"]["NORMAL"], "#14b8a6")
        self.assertEqual(payload["colors"]["CRITICAL"], "#ef4444")

    def test_tube_map_exposes_precise_indication_positions(self):
        with session_scope(DB_PATH) as session:
            result = session.scalar(
                select(IrisTubeResult)
                .join(IrisInspection, IrisInspection.iris_inspection_id == IrisTubeResult.iris_inspection_id)
                .where(IrisInspection.heat_exchanger_id == 1)
                .order_by(IrisTubeResult.iris_tube_result_id)
            )
            self.assertIsNotNone(result)
            inspection = session.get(IrisInspection, result.iris_inspection_id)
            indication = IrisIndication(
                iris_tube_result_id=result.iris_tube_result_id,
                sequence_no=9999,
                source_page=7,
                axial_position_mm=1234.5,
                circumferential_angle_deg=90.0,
                measured_thickness_mm=1.23,
                wall_loss_depth_mm=0.44,
                wall_loss_percent=26.35,
                defect_type="PIT",
                defect_length_mm=18.0,
                judgement="WARNING",
                raw_json="{}",
            )
            session.add(indication)
            session.flush()
            indication_id = indication.iris_indication_id
            tube_id = result.tube_id
            campaign_id = inspection.campaign_id
        try:
            payload = self.client.get(
                f"/api/exchangers/1/tube-map?campaign_id={campaign_id}"
            ).get_json()
            tube = next(item for item in payload["tubes"] if item["tube_id"] == tube_id)
            precise = next(
                item for item in tube["indications"]
                if item["indication_id"] == indication_id
            )
            self.assertEqual(precise["axial_position_mm"], 1234.5)
            self.assertEqual(precise["circumferential_angle_deg"], 90.0)
            self.assertEqual(precise["wall_loss_percent"], 26.35)
            self.assertEqual(precise["source_page"], 7)
            self.assertEqual(precise["position_quality"], "MEASURED")
            self.assertFalse(precise["is_representative"])
        finally:
            with session_scope(DB_PATH) as session:
                saved = session.get(IrisIndication, indication_id)
                if saved:
                    session.delete(saved)

    def test_tube_map_uses_result_location_when_indication_position_is_missing(self):
        with session_scope(DB_PATH) as session:
            result = session.scalar(
                select(IrisTubeResult)
                .join(IrisInspection, IrisInspection.iris_inspection_id == IrisTubeResult.iris_inspection_id)
                .where(IrisInspection.heat_exchanger_id == 1)
                .order_by(IrisTubeResult.iris_tube_result_id)
            )
            self.assertIsNotNone(result)
            inspection = session.get(IrisInspection, result.iris_inspection_id)
            original_location = result.location_mm
            result.location_mm = 2345.6
            indication = IrisIndication(
                iris_tube_result_id=result.iris_tube_result_id,
                sequence_no=9998,
                axial_position_mm=None,
                circumferential_angle_deg=None,
                defect_type="PIT",
                raw_json="{}",
            )
            session.add(indication)
            session.flush()
            indication_id = indication.iris_indication_id
            result_id = result.iris_tube_result_id
            tube_id = result.tube_id
            campaign_id = inspection.campaign_id
        try:
            payload = self.client.get(
                f"/api/exchangers/1/tube-map?campaign_id={campaign_id}"
            ).get_json()
            tube = next(item for item in payload["tubes"] if item["tube_id"] == tube_id)
            fallback = next(
                item for item in tube["indications"]
                if item["indication_id"] == indication_id
            )
            self.assertEqual(fallback["axial_position_mm"], 2345.6)
            self.assertEqual(fallback["position_quality"], "RESULT_REPRESENTATIVE")
            self.assertEqual(fallback["position_source"], "TUBE_RESULT")
            self.assertTrue(fallback["is_representative"])
            self.assertFalse(fallback["circumferential_position_known"])
            self.assertEqual(
                tube["measured_indication_count"],
                sum(
                    item["position_quality"] == "MEASURED"
                    for item in tube["indications"]
                ),
            )
        finally:
            with session_scope(DB_PATH) as session:
                saved = session.get(IrisIndication, indication_id)
                if saved:
                    session.delete(saved)
                saved_result = session.get(IrisTubeResult, result_id)
                saved_result.location_mm = original_location

    def test_remaining_life_formula_and_cumulative_limit(self):
        tube = calculate_tube_plugging_year(
            current_year=2026,
            nominal_thickness_mm=2.5,
            current_wall_loss_mm=0.5,
            corrosion_rate_mm_per_year=0.1,
        )
        self.assertAlmostEqual(tube["plugging_loss_mm"], 1.0)
        self.assertAlmostEqual(tube["exact_year"], 2031.0)
        self.assertEqual(tube["plugging_year"], 2031)

        earlier_threshold = calculate_tube_plugging_year(
            current_year=2026,
            nominal_thickness_mm=2.5,
            current_wall_loss_mm=0.5,
            corrosion_rate_mm_per_year=0.1,
            plugging_wall_loss_percent=30,
        )
        self.assertAlmostEqual(earlier_threshold["years_from_measurement"], 2.5)
        self.assertEqual(earlier_threshold["plugging_year"], 2028)

        summary = summarize_remaining_life(
            plugging_years=[2026, 2031, 2031],
            total_tube_count=10,
            is_critical=False,
            current_year=2026,
        )
        self.assertEqual(summary["limit_percent"], 20.0)
        self.assertEqual(summary["allowable_plugging_count"], 2)
        self.assertEqual(summary["life_reached_year"], 2031)
        self.assertEqual(summary["remaining_life_years"], 5)

        critical_default = summarize_remaining_life(
            plugging_years=[2026, 2031, 2031],
            total_tube_count=10,
            is_critical=True,
            current_year=2026,
        )
        self.assertEqual(critical_default["limit_percent"], 20.0)
        self.assertEqual(critical_default["allowable_plugging_count"], 2)
        custom_limit = summarize_remaining_life(
            plugging_years=[2031, 2032],
            total_tube_count=10,
            is_critical=False,
            current_year=2026,
            allowable_plugging_percent=15,
        )
        self.assertEqual(custom_limit["limit_percent"], 15.0)
        self.assertEqual(custom_limit["allowable_plugging_count"], 1)
        self.assertEqual(custom_limit["life_reached_year"], 2032)
        for invalid_limit in (0, 21):
            with self.assertRaisesRegex(ValueError, "1%부터 20%"):
                summarize_remaining_life(
                    plugging_years=[2031],
                    total_tube_count=10,
                    is_critical=False,
                    current_year=2026,
                    allowable_plugging_percent=invalid_limit,
                )

    def test_remaining_life_uses_measurement_date(self):
        tube = calculate_tube_plugging_year(
            measurement_date="2020-07-01",
            nominal_thickness_mm=2.5,
            current_wall_loss_mm=0.5,
            corrosion_rate_mm_per_year=0.1,
        )
        self.assertEqual(tube["projection_origin_date"], "2020-07-01")
        self.assertEqual(tube["plugging_year"], 2025)
        self.assertTrue(tube["plugging_date"].startswith("2025-"))

    def test_api510_remaining_life_and_inspection_interval(self):
        result = calculate_api510_inspection_interval(
            actual_thickness_mm=3.0,
            required_thickness_mm=2.0,
            corrosion_rate_mm_per_year=0.1,
            last_inspection_date="2026-01-01",
            inspection_type="INTERNAL_ON_STREAM",
            as_of_date="2026-08-21",
        )
        self.assertAlmostEqual(result["remaining_corrosion_allowance_mm"], 1.0)
        self.assertAlmostEqual(result["remaining_life_years"], 10.0)
        self.assertAlmostEqual(result["inspection_interval_years"], 5.0)
        self.assertEqual(result["next_inspection_date"], "2031-01-01")
        self.assertFalse(result["overdue"])

        normalized = calculate_api510_inspection_interval(
            actual_thickness_mm=3.0,
            required_thickness_mm=2.0,
            corrosion_rate_mm_per_year=0.1,
            last_inspection_date="202601-01-",
            inspection_type="INTERNAL_ON_STREAM",
            as_of_date="2026-08-21",
        )
        self.assertEqual(normalized["last_inspection_date"], "2026-01-01")
        self.assertEqual(normalized["next_inspection_date"], "2031-01-01")

        with self.assertRaisesRegex(ValueError, "YYYY-MM-DD"):
            calculate_api510_inspection_interval(
                actual_thickness_mm=3.0,
                required_thickness_mm=2.0,
                corrosion_rate_mm_per_year=0.1,
                last_inspection_date="2026-13-40",
            )

        capped = calculate_api510_inspection_interval(
            actual_thickness_mm=5.0,
            required_thickness_mm=1.0,
            corrosion_rate_mm_per_year=0.1,
            last_inspection_date="2026-01-01",
            inspection_type="EXTERNAL",
            as_of_date="2026-08-21",
        )
        self.assertEqual(capped["inspection_interval_years"], 5.0)
        with self.assertRaisesRegex(ValueError, "0보다 큰 적용 부식속도"):
            calculate_api510_inspection_interval(
                actual_thickness_mm=3.0,
                required_thickness_mm=2.0,
                corrosion_rate_mm_per_year=0,
                last_inspection_date="2026-01-01",
            )

    def test_short_long_and_governing_corrosion_rates(self):
        assessments = [
            SimpleNamespace(
                corrosion_rate_mm_per_year=0.05,
                elapsed_years=4.0,
                thickness_change_mm=0.20,
                assessed_at=datetime(2020, 1, 1),
            ),
            SimpleNamespace(
                corrosion_rate_mm_per_year=0.10,
                elapsed_years=2.0,
                thickness_change_mm=0.20,
                assessed_at=datetime(2025, 1, 1),
            ),
        ]
        result = summarize_corrosion_rates(assessments)
        self.assertAlmostEqual(result["short_term_rate_mm_per_year"], 0.10)
        self.assertAlmostEqual(result["long_term_rate_mm_per_year"], 0.4 / 6.0)
        self.assertAlmostEqual(result["selected_rate_mm_per_year"], 0.10)
        self.assertEqual(result["selection_basis"], "SHORT_TERM_GOVERNS")

    def test_asme_viii_shell_calculation_and_api510_approval_workflow(self):
        asme = calculate_asme_viii_internal_pressure(
            component_type="CYLINDRICAL_SHELL",
            inside_diameter_mm=1000,
            design_pressure_mpa=1.0,
            allowable_stress_mpa=100.0,
            joint_efficiency=1.0,
            actual_thickness_mm=10.0,
        )
        self.assertAlmostEqual(asme["required_thickness_mm"], 500 / 99.4)
        self.assertAlmostEqual(asme["mawp_mpa"], 1000 / 506)

        component_id = services.save_pressure_boundary_component(1, {
            "component_code": "TEST-SHELL-API510",
            "component_type": "CYLINDRICAL_SHELL",
            "side_code": "SHELL",
            "inside_diameter_mm": "1000",
            "design_pressure_mpa": "1.0",
            "static_head_mpa": "0",
            "allowable_stress_mpa": "100",
            "joint_efficiency": "1",
            "supplemental_thickness_mm": "0",
            "actual_thickness_mm": "10",
            "actual_thickness_date": "2026-01-01",
            "design_basis_source": "TEST CALCULATION SHEET REV.0",
            "thickness_source": "TEST CML-1",
            "code_edition": "TEST EDITION",
        })
        calculation_id = services.create_api510_calculation(component_id, {
            "inspection_type": "INTERNAL_ON_STREAM",
            "actual_thickness_mm": "10",
            "last_inspection_date": "2026-01-01",
            "short_term_corrosion_rate_mm_per_year": "0.1",
            "long_term_corrosion_rate_mm_per_year": "0.08",
            "corrosion_rate_selection_basis": "단기 부식속도 지배",
        }, actor="creator")
        self.assertTrue(services.transition_api510_calculation(
            calculation_id, "SUBMITTED", "creator"
        ))
        self.assertTrue(services.transition_api510_calculation(
            calculation_id, "APPROVED", "api-inspector", "검산 완료"
        ))
        with session_scope(DB_PATH) as session:
            calculation = session.get(Api510Calculation, calculation_id)
            component = session.get(PressureBoundaryComponent, component_id)
            self.assertEqual(calculation.status, "APPROVED")
            self.assertEqual(calculation.reviewed_by, "api-inspector")
            self.assertEqual(component.design_basis_source, "TEST CALCULATION SHEET REV.0")

    def test_asme_viii_validation_case_fixture(self):
        fixture = json.loads(
            (APP_DIR / "test_data" / "api510_validation_cases.json").read_text(encoding="utf-8")
        )
        self.assertTrue(fixture["approval_required"])
        for case in fixture["cases"]:
            result = calculate_asme_viii_internal_pressure(
                component_type=case["component_type"],
                inside_diameter_mm=case["inside_diameter_mm"],
                design_pressure_mpa=case["design_pressure_mpa"],
                allowable_stress_mpa=case["allowable_stress_mpa"],
                joint_efficiency=case["joint_efficiency"],
                actual_thickness_mm=case["actual_thickness_mm"],
            )
            self.assertAlmostEqual(
                result["required_thickness_mm"],
                case["expected_required_thickness_mm"],
                places=8,
                msg=case["case_id"],
            )
            self.assertAlmostEqual(
                result["mawp_mpa"], case["expected_mawp_mpa"], places=8,
                msg=case["case_id"],
            )

    def test_api510_management_page(self):
        response = self.client.get("/exchangers/1/api510")
        self.assertEqual(response.status_code, 200)
        self.assertIn("압력경계 부품 등록".encode("utf-8"), response.data)
        self.assertIn("API 510 계산 이력".encode("utf-8"), response.data)

    def test_same_day_readings_use_minimum_thickness_and_maximum_loss(self):
        inspection = SimpleNamespace(
            inspection_date=datetime(2025, 7, 30),
            report_date=datetime(2025, 8, 1),
        )
        readings = [
            (
                SimpleNamespace(
                    iris_tube_result_id=1,
                    minimum_wall_thickness_mm=1.80,
                    wall_loss_percent=25.0,
                ),
                inspection,
            ),
            (
                SimpleNamespace(
                    iris_tube_result_id=2,
                    minimum_wall_thickness_mm=1.85,
                    wall_loss_percent=35.0,
                ),
                inspection,
            ),
        ]
        snapshot = conservative_daily_snapshot(readings)
        self.assertEqual(snapshot["minimum_wall_thickness_mm"], 1.80)
        self.assertEqual(snapshot["maximum_wall_loss_percent"], 35.0)
        self.assertEqual(snapshot["thickness_result_id"], 1)
        self.assertEqual(snapshot["wall_loss_result_id"], 2)
        self.assertEqual(snapshot["date"].isoformat(), "2025-07-30")

    def test_measurement_location_matching_normalizes_direction(self):
        same = compare_measurement_locations(100.0, 175.0, tolerance_mm=100.0)
        self.assertTrue(same["matched"])
        self.assertEqual(same["direction_mapping"], "SAME_DIRECTION")

        reversed_position = compare_measurement_locations(
            2350.0,
            100.0,
            previous_tube_length_mm=2500.0,
            current_tube_length_mm=2450.0,
            tolerance_mm=100.0,
        )
        self.assertTrue(reversed_position["matched"])
        self.assertEqual(
            reversed_position["direction_mapping"], "REVERSE_CURRENT"
        )
        self.assertEqual(reversed_position["distance_mm"], 0.0)

        outside = compare_measurement_locations(
            2350.0,
            400.0,
            previous_tube_length_mm=2500.0,
            tolerance_mm=100.0,
        )
        self.assertFalse(outside["matched"])
        self.assertEqual(outside["reason"], "POSITION_OUTSIDE_TOLERANCE")

        missing = compare_measurement_locations(None, 100.0)
        self.assertFalse(missing["matched"])
        self.assertEqual(missing["reason"], "MISSING_LOCATION")

        raw_inspection = SimpleNamespace(raw_report_json=json.dumps({
            "report": {"iris_report_details": {
                "measurement_conditions": {"tube_length_m": 2.5}
            }}
        }))
        self.assertEqual(inspection_tube_length_mm(raw_inspection), 2500.0)

        critical = summarize_remaining_life(
            plugging_years=[2026, 2031],
            total_tube_count=10,
            is_critical=True,
            current_year=2026,
            allowable_plugging_percent=10,
        )
        self.assertEqual(critical["limit_percent"], 10.0)
        self.assertEqual(critical["allowable_plugging_count"], 1)
        self.assertEqual(critical["life_reached_year"], 2031)

    def test_remaining_life_target_year_tube_search(self):
        response = self.client.get("/remaining-life/tubes?target_year=2042&unit_code=VCM1")
        self.assertEqual(response.status_code, 200)
        self.assertIn("Tube Plugging 예측 대상 조회".encode("utf-8"), response.data)
        self.assertIn("공정".encode("utf-8"), response.data)
        self.assertIn("열교환기 코드".encode("utf-8"), response.data)
        self.assertIn("감육률".encode("utf-8"), response.data)
        self.assertIn("적용 CR".encode("utf-8"), response.data)
        self.assertIn(b'name="search"', response.data)
        self.assertIn("2042년".encode("utf-8"), response.data)

        filtered = self.client.get(
            "/remaining-life/tubes?target_year=2042&unit_code=VCM1&search=E-111"
        )
        self.assertEqual(filtered.status_code, 200)
        self.assertIn(b'value="E-111"', filtered.data)
        invalid = self.client.get("/remaining-life/tubes?target_year=1900")
        self.assertEqual(invalid.status_code, 302)

    def test_remaining_life_what_if_simulation(self):
        response = self.client.get(
            "/exchangers/1?simulate=1&simulation_limit=10&corrosion_rate_percent=150"
        )
        self.assertEqual(response.status_code, 200)
        self.assertIn("Tube Plugging 예측 시뮬레이션".encode("utf-8"), response.data)
        self.assertIn("시뮬레이션 조건".encode("utf-8"), response.data)
        self.assertIn("10% · 부식속도 150%".encode("utf-8"), response.data)

        invalid = self.client.get(
            "/exchangers/1?simulate=1&simulation_limit=21&corrosion_rate_percent=150"
        )
        self.assertEqual(invalid.status_code, 200)
        self.assertIn("시뮬레이션 입력 오류".encode("utf-8"), invalid.data)

    def test_remaining_life_dedicated_simulation_screen(self):
        page = self.client.get("/exchangers/1/remaining-life-simulation")
        self.assertEqual(page.status_code, 200)
        self.assertIn("공정 / 열교환기 코드".encode("utf-8"), page.data)
        self.assertIn("계산 방식 및 조건 설정".encode("utf-8"), page.data)
        self.assertIn(b'name="calculation_method"', page.data)
        self.assertIn(b'name="simulation_limit"', page.data)
        self.assertIn(b'type="range" name="simulation_limit"', page.data)
        self.assertIn(b'min="1" max="20" step="1"', page.data)
        self.assertIn(b'name="plugging_wall_loss_percent"', page.data)
        self.assertIn(b'name="corrosion_rate_percent"', page.data)
        self.assertIn(b'name="forecast_end_year"', page.data)

        result = self.client.get(
            "/exchangers/1/remaining-life-simulation"
            "?simulate=1&simulation_limit=10&plugging_wall_loss_percent=35"
            "&corrosion_rate_percent=150&forecast_end_year=2070"
        )
        self.assertEqual(result.status_code, 200)
        self.assertIn("사용자 Simulation".encode("utf-8"), result.data)
        self.assertIn("35.0%".encode("utf-8"), result.data)
        self.assertIn(b'"simulation"', result.data)

        invalid = self.client.get(
            "/exchangers/1/remaining-life-simulation"
            "?simulate=1&simulation_limit=10&plugging_wall_loss_percent=0"
            "&corrosion_rate_percent=150&forecast_end_year=2070"
        )
        self.assertEqual(invalid.status_code, 200)
        self.assertIn("시뮬레이션 입력 오류".encode("utf-8"), invalid.data)

        api510 = self.client.get(
            "/exchangers/1/remaining-life-simulation"
            "?simulate=1&calculation_method=API510"
            "&api_inspection_type=INTERNAL_ON_STREAM"
            "&api_actual_thickness_mm=3&api_required_thickness_mm=2"
            "&api_corrosion_rate_mm_per_year=0.1&api_last_inspection_date=2026-01-01"
            "&api_confirm_required_thickness=1"
        )
        self.assertEqual(api510.status_code, 200)
        self.assertIn("API 510 검사주기".encode("utf-8"), api510.data)
        self.assertIn("5.00년".encode("utf-8"), api510.data)
        self.assertIn("2031-01-01".encode("utf-8"), api510.data)

        malformed_date = self.client.get(
            "/exchangers/1/remaining-life-simulation"
            "?simulate=1&calculation_method=API510"
            "&api_inspection_type=INTERNAL_ON_STREAM"
            "&api_actual_thickness_mm=3&api_required_thickness_mm=2"
            "&api_corrosion_rate_mm_per_year=0.1&api_last_inspection_date=202601-01-"
            "&api_confirm_required_thickness=1"
        )
        self.assertEqual(malformed_date.status_code, 200)
        self.assertNotIn(b"Invalid isoformat string", malformed_date.data)
        self.assertIn(b'value="2026-01-01"', malformed_date.data)
        self.assertIn("2031-01-01".encode("utf-8"), malformed_date.data)

    def test_plugging_unavailable_guide_modal_for_e331d(self):
        with session_scope(DB_PATH) as session:
            exchanger_id = session.scalar(
                select(HeatExchanger.heat_exchanger_id)
                .join(Equipment, Equipment.equipment_id == HeatExchanger.equipment_id)
                .where(Equipment.equipment_code == "E-331D")
            )
        self.assertIsNotNone(exchanger_id)
        response = self.client.get(f"/exchangers/{exchanger_id}")
        self.assertEqual(response.status_code, 200)
        self.assertIn(b'data-plugging-help', response.data)
        self.assertIn(b'id="plugging-help-modal"', response.data)
        self.assertIn("Plugging 예상 산정 불가 판단 기준".encode("utf-8"), response.data)
        self.assertIn("NO_POSITIVE_CORROSION_RATE".encode("utf-8"), response.data)
        self.assertIn(b"life-prediction.js", response.data)
        tube_map = self.client.get(f"/api/exchangers/{exchanger_id}/tube-map").get_json()
        self.assertEqual(tube_map["tube_spec"]["source"], "GLOBAL_AVERAGE")
        self.assertEqual(tube_map["tube_spec"]["source_label"], "전체 열교환기 평균")
        self.assertTrue(tube_map["tube_spec"]["length_is_estimated"])
        self.assertGreater(tube_map["tube_spec"]["length_mm"], 0)
        self.assertIn("전체 열교환기 평균".encode("utf-8"), response.data)

    def test_tube_operating_status_is_separate_from_inspection_status(self):
        with session_scope(DB_PATH) as session:
            tube_id = session.scalar(select(HeatExchangerTube.tube_id).order_by(HeatExchangerTube.tube_id))
        response = self.client.post(f"/tubes/{tube_id}/status", data={
            "current_status": "PLUGGED", "reason": "튜브 폐쇄 처리",
        }, follow_redirects=True)
        self.assertEqual(response.status_code, 200)
        self.assertIn("폐쇄".encode(), response.data)
        self.assertIn("튜브 폐쇄 처리".encode(), response.data)
        with session_scope(DB_PATH) as session:
            tube = session.get(HeatExchangerTube, tube_id)
            event = session.scalar(select(TubeLifecycleEvent).where(
                TubeLifecycleEvent.tube_id == tube_id
            ).order_by(TubeLifecycleEvent.lifecycle_event_id.desc()))
            self.assertEqual(tube.current_status, "PLUGGED")
            self.assertEqual(event.event_type, "STATUS_IN_SERVICE_TO_PLUGGED")

        payload = self.client.get("/api/exchangers/1/tube-map").get_json()
        mapped = next(item for item in payload["tubes"] if item["tube_id"] == tube_id)
        self.assertEqual(mapped["operational_status"], "PLUGGED")
        self.assertIn(mapped["inspection_status"], ("INSPECTED", "NOT_INSPECTED"))
        self.assertGreaterEqual(payload["summary"]["operational"]["PLUGGED"], 1)


if __name__ == "__main__":
    unittest.main()
